"""Tests for rag_store (library ingest + hybrid search). Run: .venv/bin/python -m pytest -q test_rag_store.py"""
import csv
import datetime as dt
import json
import os
import sys
import pathlib

import pytest

HERE = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))
os.environ.setdefault("EMBED_THREADS", "2")

SECTIONS = [
    ("1 The Heart", 1, [
        "The heart is a muscular organ that pumps blood through the circulatory system.",
        "It has four chambers: two atria and two ventricles.",
        "The sinoatrial node in the right atrium acts as the natural pacemaker and sets the heart rate at about seventy beats per minute.",
        "Cardiac output is the product of heart rate and stroke volume.",
    ]),
    ("1.1 Heart Valves", 2, [
        "The mitral valve separates the left atrium from the left ventricle.",
        "The aortic valve opens when left ventricular pressure exceeds aortic pressure during systole.",
        "Valve disease such as stenosis narrows the opening and increases the work of the ventricle.",
    ]),
    ("2 The Kidney", 1, [
        "The kidney filters about one hundred and eighty litres of plasma each day in the glomeruli.",
        "The loop of Henle creates a concentration gradient in the medulla that allows the body to produce concentrated urine.",
        "Furosemide is a loop diuretic that blocks the sodium potassium chloride cotransporter in the thick ascending limb.",
    ]),
    ("2.1 The Cheek", 2, [
        "The zygomaticofacial nerve emerges through a small foramen on the zygomatic bone and supplies the skin over the cheek.",
        "Sensation from the face is carried mostly by the trigeminal nerve and its three divisions.",
    ]),
    ("References", 1, [
        "Smith J, Jones K. Pacemaker cells of the heart and kidney. J Physiol. 1999;12:1-10.",
        "Brown A. Loop diuretics and the heart rate. Kidney Int. 2004;3:44-50.",
    ]),
]


def make_pdf(path: pathlib.Path) -> None:
    from reportlab.lib.pagesizes import A5
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A5)
    W, H = A5
    c.addPageLabel(0, style="D", start=101)  # printed page labels start at 101
    for pi, (title, lvl, paras) in enumerate(SECTIONS):
        c.setFont("Helvetica", 8)
        c.drawString(40, H - 30, "Synthetic Physiology Notes for Testing")  # running header
        c.drawString(W / 2, 25, str(101 + pi))                               # footer page number
        c.bookmarkPage(f"s{pi}")
        c.addOutlineEntry(title, f"s{pi}", level=lvl - 1)
        y = H - 70
        c.setFont("Helvetica-Bold", 16 if lvl == 1 else 13)
        c.drawString(40, y, title)
        y -= 30
        c.setFont("Helvetica", 10)
        for para in paras:
            words = para.split()
            line = ""
            lines = []
            for w in words:
                if len(line) + len(w) + 1 > 58:
                    lines.append(line)
                    line = w
                else:
                    line = (line + " " + w).strip()
            lines.append(line)
            # force one hyphenated line break on the Henle sentence
            if "concentration gradient" in para:
                lines = ["The loop of Henle creates a concen-", "tration gradient in the medulla that allows the body to",
                         "produce concentrated urine."]
            for ln in lines:
                c.drawString(40, y, ln)
                y -= 13
            y -= 8
        c.showPage()
    c.save()


MD = """# Chapter 7 Microbes

Intro to microbes and how they grow in culture.

## 7.1 Bacterial Shapes

Cocci are spherical bacteria. Bacilli are rod shaped bacteria. Spirilla are spiral shaped bacteria that move with flagella.

## 7.2 Gram Staining

The Gram stain uses crystal violet and iodine. Gram positive bacteria keep the purple dye because of their thick peptidoglycan wall.
Gram negative bacteria lose the dye during alcohol washing and are counterstained pink with safranin.
"""


def make_xlsx(path: pathlib.Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["order_id", "order_date", "channel", "customer", "city", "sku", "product", "size", "qty", "amount", "status", "courier", "tracking"])
    chans = ["Shop", "Website", "Marketplace"]
    stats = ["shipped", "delivered", "returned"]
    for i in range(30):
        ws.append([f"SMT-{1020 + i}", dt.datetime(2026, 8 + i % 2, 1 + i), chans[i % 3], f"Customer{i}", ["Springfield", "Riverton", "Lakeside"][i % 3],
                   f"SMT-FJ-0{i % 5}", ["Rain jacket", "Shirt"][i % 2], ["S", "M", "L"][i % 3], 1 + i % 2, 1999 + 500 * (i % 3),
                   stats[i % 3], "CourierA", f"DLV{900000 + i}"])
    inv = wb.create_sheet("Inventory")
    inv.append(["sku", "product", "size", "colour", "stock"])
    stock = {("Rain jacket", "S"): 5, ("Rain jacket", "M"): 9, ("Rain jacket", "L"): 4, ("Shirt", "S"): 12, ("Shirt", "M"): 7, ("Shirt", "L"): 2}
    for (p, s), n in stock.items():
        inv.append([f"SMT-{p[:2].upper()}-{s}", p, s, "navy", n])
    ret = wb.create_sheet("Returns")
    ret.append(["return_id", "order_id", "return_date", "reason", "refund"])
    for i in range(8):
        ret.append([f"RET-{500 + i}", f"SMT-{1020 + i}", dt.datetime(2026, 9 if i < 5 else 10, 3 + i), ["size issue", "defect"][i % 2], 1999])
    wb.save(str(path))


@pytest.fixture(scope="module")
def lib(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("rag")
    import rag_store
    rag_store.MEMORY_DIR = tmp
    corpus = tmp / "docs" / "study"
    corpus.mkdir(parents=True)
    make_pdf(corpus / "physio.pdf")
    (corpus / "microbes.md").write_text(MD)
    (corpus / "README.md").write_text("# Put your books here\n\nThis is a readme and must be skipped.")
    (corpus / "manifest.json").write_text(json.dumps([
        {"file": "physio.pdf", "title": "Synthetic Physiology Notes", "short_title": "Physio Notes", "subject": "Physiology", "kind": "book", "licence": "CC0"},
        {"file": "microbes.md", "title": "Microbiology Chapter 7", "short_title": "Micro Ch7", "subject": "Microbiology", "kind": "chapter", "licence": "CC0"},
    ]))
    shop = tmp / "docs" / "shop"
    shop.mkdir(parents=True)
    make_xlsx(shop / "Orders_FY26.xlsx")
    with open(shop / "Orders_FY26_Orders.csv", "w", newline="") as f:  # duplicate export of the Orders sheet → deduped
        csv.writer(f).writerows([["order_id", "status"], ["SMT-1020", "shipped"], ["SMT-1021", "delivered"]])
    L = rag_store.Library("testlib")
    events = []
    res = L.ingest_dir(corpus, progress=events.append)
    L._events, L._res, L._shop = events, res, shop
    return L


def test_ingest_report(lib):
    r = lib._res
    assert r["docs"] == 2 and not r["errors"], r
    assert all("README" not in f["file"] for f in r["files"])
    d = {x["file"]: x for x in lib.docs()}
    assert d["physio.pdf"]["short_title"] == "Physio Notes" and d["physio.pdf"]["pages"] == 5
    assert d["physio.pdf"]["licence"] == "CC0" and d["microbes.md"]["subject"] == "Microbiology"
    assert d["physio.pdf"]["report"]["headings_source"] == "outline"


def test_progress_events(lib):
    phases = [e["phase"] for e in lib._events]
    for p in ("extract", "chunk", "embed", "index", "done"):
        assert p in phases, phases
    e = next(e for e in lib._events if e["phase"] == "extract")
    assert {"file", "page", "pages", "chunks", "per_s", "eta_s"} <= set(e)


def test_page_label_and_section(lib):
    r = lib.search("What is the natural pacemaker of the heart?", k=3)
    p = r["passages"][0]
    assert p["page_label"] == "101" and "The Heart" in p["section"], p
    assert r["gate"] is True
    assert r["context"].startswith("[Physio Notes, p. 101, 1 The Heart]"), r["context"]
    r = lib.search("which valve separates the left atrium from the left ventricle", k=3)
    assert r["passages"][0]["page_label"] == "102" and r["passages"][0]["section"].endswith("1.1 Heart Valves")


def test_header_footer_stripped_and_dehyphenated(lib):
    rows = lib.db.execute("SELECT text FROM chunks").fetchall()
    allt = " ".join(r[0] for r in rows)
    assert "Synthetic Physiology Notes for Testing" not in allt
    assert "concentration gradient" in allt and "concen-" not in allt


def test_fts_rare_term(lib):
    r = lib.search("zygomaticofacial", k=3)
    assert r["hits"]["fts"] >= 1
    assert r["passages"][0]["page_label"] == "104" and "Cheek" in r["passages"][0]["section"]


def test_references_not_indexed(lib):
    r = lib.search("Smith Jones pacemaker cells J Physiol 1999", k=3)
    assert all("References" not in p["section"] for p in r["passages"])


def test_markdown_section(lib):
    r = lib.search("why do gram positive bacteria stay purple", k=3)
    p = r["passages"][0]
    assert p["short_title"] == "Micro Ch7" and "Gram Staining" in p["section"]
    assert r["gate"]


def test_gate_false_for_unrelated(lib):
    for q in ("Who won the football world cup in 2014?", "What's a good recipe for chocolate cake?", "How do I change a car tyre?"):
        r = lib.search(q, k=3)
        assert r["gate"] is False, (q, r["top_score"])


@pytest.mark.parametrize("budget", [40, 80, 160])
def test_budget(lib, budget):
    import embedder
    r = lib.search("How does the kidney concentrate urine and what does furosemide block?", k=3, budget_tokens=budget)
    n = len(embedder.get().tok.encode(r["context"], add_special_tokens=False).ids)
    assert 0 < n <= budget, (n, r["context"])


def test_subject_filter(lib):
    r = lib.search("bacteria", k=3, subject="Physiology")
    assert all(p["subject"] == "Physiology" for p in r["passages"])


def test_skip_duplicate(lib):
    before = lib.stats()["chunks"]
    r = lib.ingest_file(lib._res["files"][0]["file"] and next(pathlib.Path(lib.path).parent.glob("docs/study/physio.pdf")))
    assert r["skipped"] is True and lib.stats()["chunks"] == before


def test_latency(lib):
    lib.search("warm up query", k=3)
    r = lib.search("what does furosemide block", k=3)
    assert r["ms"]["total"] < 100


# ---- spreadsheets ----

@pytest.fixture(scope="module")
def shop(lib):
    res = lib.ingest_dir(lib._shop)
    return res


def test_sheet_ingest_dedupes_csv(shop):
    assert shop["docs"] == 1 and shop["deduped_csv"] == ["Orders_FY26_Orders.csv"], shop


def test_exact_order_id(lib, shop):
    for q in ("where is order SMT-1042?", "where is order SMT 1042", "status of smt-1042",
              "where is order s and t ten forty two", "where's order one thousand forty two"):
        r = lib.search(q, k=3)
        p = r["passages"][0]
        assert "SMT-1042" in p["text"] and p["page_label"] == "Orders row 24", (q, p)
        assert r["gate"] and r["context"].startswith("[Orders FY26, Orders sheet, row 24]"), r["context"]
    r = lib.search("which order shipped with tracking number DLV900007", k=3)
    assert "SMT-1027" in r["passages"][0]["text"]


def test_stock_question(lib, shop):
    r = lib.search("how many Rain jackets in size L are in stock?", k=3)
    assert "stock: 4" in r["context"] or "L 4" in r["context"], r["context"]
    assert r["gate"]


def test_month_aggregate(lib, shop):
    r = lib.search("what were the returns in September?", k=3)
    assert "September 2026" in r["context"] and "5 returns" in r["context"], r["context"]


def test_spoken_id_normaliser():
    import rag_store
    sch = {"prefixes": {"SMT": (4, (1000, 1999)), "RET": (3, (300, 399))}, "bare": ("SMT", (1000, 1999)),
           "sku_prefix": "SMT", "sku_cats": ["BG", "FJ"]}
    assert rag_store.normalize_query("Where is order S, and T, one thousand forty two?", sch) == "Where is order SMT-1042?"
    assert rag_store.normalize_query("has S M T eleven oh five shipped", sch) == "has SMT-1105 shipped"
    assert rag_store.normalize_query("How many smtf j01 do we have?", sch) == "How many SMT-FJ01 do we have?"
    assert rag_store.normalize_query("I sold 1042 jackets", sch) == "I sold 1042 jackets"
    learned = rag_store.learn_id_scheme(["Orders row 3: order id: SMT-%d · sku: SMT-FJ0%d" % (1040 + i, i % 4 + 1) for i in range(20)])
    assert learned["prefixes"]["SMT"] == (4, (1000, 1999)) and learned["sku_cats"] == ["FJ"]


def test_context_never_empty_for_long_row(lib, shop):
    r = lib.search("where is order SMT-1042", k=3, budget_tokens=30)
    assert r["context"] and r["tokens"] <= 30 and "SMT-1042" in r["context"], r


def test_spreadsheet_kind_citation(tmp_path):
    import rag_store
    L = rag_store.Library("kindtest", path=tmp_path / "k.sqlite")
    make_xlsx(tmp_path / "w.xlsx")
    L.ingest_file(tmp_path / "w.xlsx", meta={"title": "Orders workbook", "short_title": "Orders workbook", "kind": "spreadsheet"})
    r = L.search("where is order SMT-1030", k=3)
    assert r["context"].startswith("[Orders workbook, Orders sheet, row 12]"), r["context"]
    assert "p. " not in r["context"]


def test_twins_skip_and_short_titles(tmp_path):
    import rag_store
    from reportlab.pdfgen import canvas
    d = tmp_path / "c"
    d.mkdir()
    md = "# Diary\n\n## Tests\n\nMaths unit test on Tuesday 27 October covers tables of seven.\n\n## Homework\n\nScience homework: draw three plants that grow in water.\n"
    (d / "diary.md").write_text(md)
    c = canvas.Canvas(str(d / "diary.pdf"))
    c.drawString(50, 800, "Diary")
    c.showPage()
    c.drawString(50, 800, "Tests")
    c.drawString(50, 780, "Maths unit test on Tuesday 27 October covers tables of seven.")
    c.showPage()
    c.drawString(50, 800, "Homework")
    c.drawString(50, 780, "Science homework: draw three plants that grow in water.")
    c.save()
    (d / "Old-Story-English.md").write_text("# Old story\n\nThe river was poisoned.")
    (d / "manifest.json").write_text(json.dumps([
        {"file": "diary.pdf", "title": "School diary", "short_title": "School diary", "kind": "diary"},
        {"file": "diary.md", "title": "School diary (text)", "short_title": "School diary (text)", "kind": "diary"},
        {"file": "Old-Story-English.md", "exclude": True}]))
    L = rag_store.Library("kid", path=tmp_path / "r.sqlite")
    res = L.ingest_dir(d)
    assert res["deduped_twins"] == ["diary.pdf"] and res["docs"] == 1, res
    assert any("excluded" in (f.get("reason") or "") for f in res["files"])
    assert [x["short_title"] for x in L.docs()] == ["School diary"]
    r = L.search("what homework do I have for science", k=3)
    assert r["gate"] and r["passages"][0]["page_label"] == "3", r
    assert rag_store.spoken_short_title({"short_title": "Smith 2024 (100 cases)"}, "x", "f") == "Smith 2024"
    assert rag_store.spoken_short_title({"short_title": "The Clever Crow and the Big Troubl"}, "The Clever Crow and the Big Trouble (Example Press)", "f") == "The Clever Crow and the Big Trouble"
    assert rag_store.spoken_short_title({"short_title": "Reader front matter", "chapter": "Unit 1 and Ch 1 Colours"}, "t", "f") == "Reader Ch 1"


def test_reload_after_external_ingest(tmp_path):
    import rag_store
    a = rag_store.Library("ext", path=tmp_path / "e.sqlite")
    b = rag_store.Library("ext", path=tmp_path / "e.sqlite")   # e.g. the server's cached instance
    (tmp_path / "n.md").write_text("# Notes\n\nThe quokka is a small marsupial found on Rottnest Island near Perth.")
    a.ingest_file(tmp_path / "n.md")
    r = b.search("where do quokkas live", k=1)
    assert r["passages"] and "Rottnest" in r["passages"][0]["text"]
