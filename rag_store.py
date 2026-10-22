"""Local document library for Nova: ingest PDFs / Markdown / text / spreadsheets, hybrid search.

- One SQLite file per profile: data/<profile>_library.sqlite (separate from the facts DB).
- Ingest: pypdfium2 text per page + printed page labels; repeated header/footer and page-number
  lines dropped; de-hyphenation; headings from the PDF outline, else font-size/bold heuristics
  (Markdown: # headings); reference lists are stored but not indexed (kind='refs').
  Chunks never cross a section, are sentence-aligned and ≈ 120–240 wordpieces, and carry a
  heading path ("Short title › Chapter › Section") used for embedding and FTS.
  Spreadsheets (.xlsx/.csv): one chunk per row ("Orders row 42: order_id: ORD-1042 · …") plus
  per-sheet summary chunks (counts/sums by category, by month, category × size pivots).
- Index: FTS5 BM25 (header, text, codes) where `codes` holds hyphen-free exact IDs/SKUs
  ("ORD-1042" → "ord1042"); chunk vectors (MiniLM int8 + scale in SQLite, float32 matrix in RAM,
  exact scan); sentence vectors in SQLite, loaded only for the top chunks at query time.
- Search: FTS top-30 ∥ vector top-30 → RRF(k=60) → sentence-level extractive compression to a
  token budget → gate (confident enough to answer from the notes?).
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import logging
import os
import re
import sqlite3
import statistics
import threading
import time
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Callable

import numpy as np

import embedder

log = logging.getLogger("nova.rag")

HERE = Path(__file__).resolve().parent
MEMORY_DIR = Path(os.environ.get("MEMORY_DIR", HERE / "data"))
CHUNK_MAX_TOK = int(os.environ.get("RAG_CHUNK_MAX_TOK", "240"))   # wordpieces of body text
CHUNK_MIN_TOK = 80                                                 # merge smaller tail into previous
RRF_K = 60
WINDOW = (1, -1)
DOC_MENTION_BOOST = 2.0    # query names a document (author surname, distinctive title word)
PAPER_FRONT_BOOST = 1.15   # abstract / introduction / first page of a paper
CAND = 30
GROUP_BY_BOOST = 2.0  # aggregate question naming the summary's group-by column ("which city …")
LEAD_MARGIN = 0.08  # sentence score margin for leading the context with a passage other than the top-ranked one
GATE_COS = float(os.environ.get("RAG_GATE_COS", "0.45"))          # best sentence cosine
# Gate (calibrated on question sets over textbooks, papers, stories and spreadsheets): open when the best sentence cosine
# clears tau AND enough of the query's content words occur in a retrieved chunk. tau is lower for notes, diaries,
# spreadsheets and business docs (short rows embed with lower cosine). A content word that occurs nowhere in the
# library ("followers", "cricket") needs a much stronger dense match.
GATE_LENIENT_DELTA = 0.07
GATE_STRONG = 0.15
GATE_LEX_BAND = 0.10
GATE_OOV_EXTRA = 0.20
GATE_COVER = 0.6            # share of content words that must occur in a retrieved chunk (2-word questions: both)
GATE_MIN_WORDS = 2          # "the lateral one" (a quiz reply) is not a library question unless it names a rare term
GATE_RARE_DF = 12           # a content word in ≤ 12 chunks counts as rare
LENIENT_KINDS = {"notes", "diary", "spreadsheet", "sheet", "business_doc", "chapter"}
# Aggregate questions ("top customer", "best sellers last quarter", "how many … in stock", "which … below reorder")
# prefer the per-sheet summary chunks; a few spoken words map onto the words the summaries use.
_AGG_Q_RX = re.compile(r"\b(top|best|most|least|total|overall|average|which (products|items|orders|skus|customers)"
                       r"|list|below|above|per month|by month|last quarter|this quarter|quarter|why)\b", re.I)
SUMMARY_BOOST = 1.25
_EXPAND = {"why": ["reason"], "best": ["top"], "sellers": ["top", "sellers"], "bestsellers": ["top", "sellers"], "sold": ["sellers", "units"],
           "supplies": ["supplier"], "supply": ["supplier"], "returning": ["returns", "reason"], "stock": ["stock", "inventory"]}
_ANAT = {"supplies": ["supplied", "nerves"], "supply": ["supplied", "nerves"], "innervates": ["supplied"],
         "innervation": ["supplied", "nerves"], "innervated": ["supplied"]}


def _latin(w: str) -> list[str]:
    """Some texts use the Latin form of anatomical terms ending in -oid (-oideus): add that form as a search term."""
    out = list(_ANAT.get(w, []))
    if len(w) > 5 and w.endswith("oid"):
        out.append(w + "eus")
    elif len(w) > 6 and w.endswith("oids"):
        out.append(w[:-1] + "eus")
    return out


_GATE_IGNORE = set("""yesterday today tomorrow tonight doing done happen happened happening ever really called name thing things stuff
lot lots good bad get got make made go going come came tell said say want need like time way kind sort right now still yet
again also much many more most else something anything everything someone anyone one ones two first last goes bedtime
story stories""".split())
# Diaries, planners and notes answer schedule questions ("when is my maths test?") whose words are common in textbooks.
_SCHEDULE_Q_RX = re.compile(r"\b(test|tests|exam|exams|homework|due|deadline|timetable|schedule|tomorrow|today|tonight|this week|next week"
                            r"|monday|tuesday|wednesday|thursday|friday|saturday|sunday|when is|when's|what time|holiday|break|meeting)\b", re.I)
SCHEDULE_KINDS = {"diary", "notes", "planner", "calendar", "schedule"}
SCHEDULE_BOOST = 1.4
# Books do not hold recent events: "who won the match yesterday", "latest news", "price today" need a strong match
# unless the hit is the person's own notes/diary/sheet.
_RECENT_Q_RX = re.compile(r"\b(yesterday|today|tonight|latest|news|right now|currently|this (week|month|year)|score|price)\b", re.I)
SENT_MIN_WORDS = 5
SCHEMA_VER = 1
CHUNKER_VER = 1
SUPPORTED = {".pdf", ".md", ".markdown", ".txt", ".xlsx", ".csv"}

_STOP = set("""a an the and or but if of to in on at by for with from into onto over under about as is are was were be been
being do does did doing have has had having i me my we our you your he she it its they them their this that these those
what which who whom whose when where why how there here can could would should will shall may might must not no yes
please tell explain describe give show say said says about any some all more most much many very also than then so such
just only own same other again further once up down out off both each few nor too s t don now let lets know want
would like does doesn whats what's ok okay hey nova according book books notes""".split())

_ABBR = {"e.g", "i.e", "etc", "fig", "figs", "et al", "al", "dr", "mr", "mrs", "ms", "vs", "approx", "no", "nos", "vol",
         "eq", "eqs", "ref", "refs", "sec", "ch", "pp", "p", "st", "spp", "sp", "mg", "ml", "cf", "resp", "ca", "inc", "ltd",
         "jr", "sr", "prof", "dept", "u.s", "e.coli", "tab", "max", "min", "ave", "co"}
_SENT_SPLIT = re.compile(r"(?<=[.!?])[\"')\]]?\s+(?=[\"'(\[]?[A-Z0-9])")
_WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'’]*")
_CODE = re.compile(r"\b[A-Za-z0-9]*\d[A-Za-z0-9]*(?:[-/][A-Za-z0-9]+)+\b|\b[A-Za-z]+(?:[-/][A-Za-z0-9]+)*[-/][A-Za-z]*\d[A-Za-z0-9]*\b|\b(?=[A-Za-z]*\d)(?=\d*[A-Za-z])[A-Za-z0-9]{5,}\b")
_SPOKEN_CODE = re.compile(r"\b([A-Za-z]{2,5})[\s-]+(\d{2,6})\b")
_REFS_RX = re.compile(r"^\s*(\d+(\.\d+)*\.?\s+)?(references|bibliography|works cited|literature cited|reference list|cited literature|further reading|suggested readings?)\s*:?\s*$", re.I)
_BACKMATTER_RX = re.compile(r"^\s*(\d+(\.\d+)*\.?\s+)?(acknowledge?ments?|funding( information| sources)?|competing (financial )?interests?|conflicts? of interests?|declaration of (competing )?interests?|author contributions?|contributions|disclosures?|data availability( statement)?|code availability|ethics (declarations?|statement)|additional information|supplementary (information|materials?)|reporting summary|peer review( information)?|author information|correspondence)\b\s*[.:]?\s*(?P<rest>.*)$", re.I)
_BOILER_RX = re.compile(r"this book was made possible|creative commons|attribution 4\.0|attribution 3\.0|about this story|thank you for buying|story attribution|illustration attribution|images attributions|disclaimer:|isbn|printed (in|at)|reprint 20|first edition|all rights reserved", re.I)
_EXERCISE_RX = re.compile(r"^\s*(\d+(\.\d+)*\.?\s+)?(review questions|critical thinking questions|interactive link questions|check your understanding|exercises|problems|self[- ]test)\s*$", re.I)
_PAGENUM_RX = re.compile(r"^\s*(page\s+)?(\d{1,4}|[ivxlcdm]{1,7})(\s*(of|/)\s*\d{1,4})?\s*$", re.I)
_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]

ProgressFn = Callable[[dict], None] | None


# ───────────────────────────── text helpers ─────────────────────────────

def _norm_space(s: str) -> str:
    return " ".join(s.split())


def _clean_text(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("\x02", "").replace("\xad", "").replace("\ufffe", "").replace("\u00a0", " ")  # PDFium marks line-end hyphens as \x02 and joins the lines
    s = s.replace("ﬁ", "fi").replace("ﬂ", "fl")
    return s


_OCR_U = re.compile(r"\b([A-Za-z]?[a-z]+)U([a-z]+)\b")


def _repair_ocr_l(pages_lines: list) -> list:
    """Old scans OCR "li"/"ll" as "U" ("suppUed", "ceUs"). Pick the spelling that occurs elsewhere in the same
    document; leave the word alone when neither does."""
    import collections
    vocab = collections.Counter(w.lower() for lines in pages_lines for l in lines for w in re.findall(r"[A-Za-z]+", l[0]))

    def fix(m: re.Match) -> str:
        a, b = m.group(1), m.group(2)
        best = max(("li", "ll"), key=lambda x: vocab[(a + x + b).lower()])
        return a + best + b if vocab[(a + best + b).lower()] else m.group(0)

    return [[(_OCR_U.sub(fix, l[0]),) + tuple(l[1:]) for l in lines] for lines in pages_lines]


def split_sentences(text: str) -> list[str]:
    text = _norm_space(text)
    if not text:
        return []
    parts = _SENT_SPLIT.split(text)
    out: list[str] = []
    for p in parts:
        if out:
            prev = out[-1]
            m = re.search(r"(\b[\w.]+)\.$", prev)
            last = (m.group(1).lower() if m else "")
            if last in _ABBR or re.fullmatch(r"[A-Z]", prev[-2:-1] or "x") and len(prev) > 1 and prev[-3:-2] in (" ", ""):
                out[-1] = prev + " " + p
                continue
        out.append(p)
    final: list[str] = []
    for s in out:  # very long "sentences" (tables, run-ons) → word windows of 60
        w = s.split()
        if len(w) > 90:
            for i in range(0, len(w), 60):
                final.append(" ".join(w[i:i + 60]))
        else:
            final.append(s)
    return final


def _terms(q: str) -> list[str]:
    seen, out = set(), []
    for w in _WORD.findall(q.lower()):
        w = w.replace("’", "'").split("'")[0]
        if len(w) < 2 or w in _STOP or w in seen:
            continue
        seen.add(w)
        out.append(w)
    return out


def _codes(text: str) -> list[str]:
    out = []
    for m in _CODE.finditer(text):
        c = re.sub(r"[-/]", "", m.group(0)).lower()
        if len(c) >= 4 and re.search(r"\d", c):
            out.append(c)
    return out


def _query_codes(q: str) -> list[str]:
    cs = set(_codes(q))
    for m in _SPOKEN_CODE.finditer(q):
        if m.group(1).lower() not in _STOP:
            cs.add((m.group(1) + m.group(2)).lower())
    return sorted(cs)


def _stem5(w: str) -> str:
    """Crude stem for overlap tests: drop plural/verb endings, keep 5 chars ("peels" → "peel", "dustbins" → "dustb")."""
    if len(w) > 4:
        w = re.sub(r"(ies|es|s|ing|ed)$", "", w)
    return w[:5]


def _quant(V: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    scale = np.maximum(np.abs(V).max(1), 1e-9) / 127.0
    return np.round(V / scale[:, None]).astype(np.int8), scale.astype(np.float32)


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def find_manifest_meta(path: Path) -> dict:
    """Look for manifest.json next to the file or in up to 3 parent folders; match by file name."""
    path = Path(path).resolve()
    d = path.parent
    for _ in range(4):
        mf = d / "manifest.json"
        if mf.exists():
            try:
                data = json.loads(mf.read_text())
            except Exception as e:  # noqa: BLE001
                log.warning("bad manifest %s: %s", mf, e)
                data = None
            items: list = []
            if isinstance(data, list):
                items = data
            elif isinstance(data, dict):
                for v in data.values():
                    if isinstance(v, list) and v and isinstance(v[0], dict):
                        items.extend(v)
                if "file" in data:
                    items.append(data)
            rel = str(path.relative_to(d)) if path.is_relative_to(d) else path.name
            for it in items:
                if not isinstance(it, dict):
                    continue
                f = str(it.get("file", ""))
                if f and (f == rel or f == path.name or Path(f).name == path.name):
                    return dict(it)
        if d.parent == d:
            break
        d = d.parent
    return {}


# Files never indexed for a profile (corpus files are left untouched). Manifest entries can also say
# "exclude": true / "index": false. Library.exclude(pattern) adds patterns at runtime (stored in the DB).
DEFAULT_SKIP: dict[str, list[str]] = {
}


def spoken_short_title(m: dict, title: str, fallback: str) -> str:
    """Short title for speech/citations: no parentheses, no mid-word truncation, chapter-aware for book chapters."""
    short = (m.get("short_title") or "").strip()
    chapter = str(m.get("chapter") or "")
    front_rx = r"\b(?:prelims|front matter)\b"  # a book's front-matter file named after the whole book
    if short and re.search(front_rx, short, re.I) and chapter and not re.match(r"\s*(?:prelims|front matter)", chapter, re.I):
        n = re.search(r"\bCh(?:apter)?\.?\s*(\d+)", chapter)
        base = re.sub(r"\s*" + front_rx, "", short, flags=re.I).strip()
        short = f"{base} Ch {n.group(1)}" if n else base
    if not short:
        short = title if len(title) <= 40 else " ".join(title.split()[:5])
    bare = re.sub(r"\s*\([^)]*\)", "", title).strip()
    if bare.startswith(short) and short != bare and bare[len(short):len(short) + 1].isalnum():
        short = bare if len(bare) <= 45 else short.rsplit(" ", 1)[0]   # a title cut mid-word → full title
    # "(Author 2024)" / "(Abbr)" → ", Author 2024"; descriptive asides "(100 cases)", "(CSV)" are dropped
    short = re.sub(r"\s*\(([A-Z][a-z]+(?: \d{4})?)\)", r", \1", short)
    short = re.sub(r"\s*[\(\[][^\)\]]*[\)\]]", "", short).strip(" -–:,") or fallback
    return short


def _twin_pdf(path: Path) -> Path | None:
    for ext in (".pdf", ".PDF"):
        q = path.with_suffix(ext)
        if q.exists():
            return q
    return None


def _map_blocks_to_pdf(blocks: list[tuple], pdf: Path) -> tuple[list[tuple], list[str]] | None:
    """Give Markdown blocks the page of the twin PDF where their text appears (for 'p. N' citations)."""
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(pdf))
        pages, labels = [], []
        for i in range(len(doc)):
            tp = doc[i].get_textpage()
            pages.append(re.sub(r"[^a-z0-9]+", "", _clean_text(tp.get_text_range()).lower()))
            lab = ""
            try:
                lab = doc.get_page_label(i).strip()
            except Exception:  # noqa: BLE001
                pass
            labels.append(lab or str(i + 1))
            tp.close()
        doc.close()
    except Exception as e:  # noqa: BLE001
        log.warning("twin pdf %s unreadable: %s", pdf.name, e)
        return None
    cur, out, found = 0, [], 0
    for kind, lvl, text, _pg in blocks:
        key = re.sub(r"[^a-z0-9]+", "", text.lower())[:40]
        if len(key) >= 12:
            for j in list(range(cur, len(pages))) + list(range(0, cur)):
                if key in pages[j]:
                    cur = j
                    found += 1
                    break
        out.append((kind, lvl, text, cur))
    return (out, labels) if found else None


# ───────────────────────────── extraction ─────────────────────────────
# A document is turned into a list of blocks: ("h", level, text, page_idx) or ("t", 0, text, page_idx).

def _line_heights(textpage, raw, text: str) -> list[tuple[str, float, int, int]]:
    """Split page text into lines; per line: (text, loose box height, max weight, char offset)."""
    out = []
    off = 0
    n = textpage.count_chars()
    rect = raw.FS_RECTF()
    for line in text.split("\n"):
        L = len(line) + 1
        s = line.rstrip("\r")
        st = len(s) - len(s.lstrip())
        idxs = [off + st + k for k in (0, len(s.strip()) // 2, max(len(s.strip()) - 1, 0))]
        hs, ws = [], []
        for i in idxs:
            if 0 <= i < n and raw.FPDFText_GetLooseCharBox(textpage, i, rect):
                hs.append(abs(rect.top - rect.bottom))
                ws.append(raw.FPDFText_GetFontWeight(textpage, i))
        out.append((s, (sorted(hs)[len(hs) // 2] if hs else 0.0), (max(ws) if ws else 0), off))
        off += L
    return out


def _special(title: str, lvl):
    """Level 9 = reference list, 8 = exercise/review questions; both stored but not indexed."""
    if _REFS_RX.match(title):
        return 9
    m = _BACKMATTER_RX.match(title)
    if m and not m.group("rest").strip(" .:"):
        return 9
    if _EXERCISE_RX.match(title) or _FRONT_RX.fullmatch(title.strip()):
        return 8
    return lvl


_FRONT_RX = re.compile(r"(\d+\.?\s*)?((table of )?contents|index|glossary of terms|prefaces?( to the [\w ]+)?|foreword|about the authors?"
                       r"|about (this|the) (book|series|edition)|how to use this book|notes? to the (reader|teacher)s?|dedication)\.?", re.I)


def _heading_like(s: str) -> bool:
    """Reject OCR debris as headings: mostly letters, at least one real word."""
    ns = re.sub(r"\s", "", s)
    return (len(ns) >= 4 and sum(ch.isalpha() for ch in ns) >= 0.75 * len(ns)
            and bool(re.search(r"[A-Za-z]{4,}", s)) and not re.search(r"[\^|~=<>{}\\]", s))


def _printed_labels(pages_lines: list, pdf_labels: list[str]) -> tuple[list[str], str]:
    """Printed page numbers from header/footer lines ("212 CHAPTER TITLE", "758 Chapter 18 | Blood", "17").

    A number counts only when its offset to the page index recurs across the document. If the PDF's own
    labels agree with the printed numbers (or nothing is detected), the PDF labels are kept; otherwise
    the printed numbers are used, interpolated across ≤ 6 pages without a detected number."""
    n = len(pages_lines)
    cand: dict[int, set[int]] = {}
    for i, lines in enumerate(pages_lines):
        ne = [l[0].strip() for l in lines if l[0].strip()]
        cs = set()
        for s in ne[:2] + ne[-2:]:
            if len(s.split()) > 14:
                continue
            for m in (re.match(r"^(\d{1,4})\b", s), re.search(r"\b(\d{1,4})$", s)):
                if m and not re.match(r"^\d+(\.\d+)+", s):
                    cs.add(int(m.group(1)))
        if cs:
            cand[i] = cs
    offs = Counter(c - i for i, cs in cand.items() for c in cs)
    good = {o for o, k in offs.items() if k >= max(3, 0.02 * n)}
    printed: dict[int, int] = {}
    for i, cs in cand.items():
        hits = [c for c in cs if c - i in good]
        if len(hits) == 1:
            printed[i] = hits[0]
    if len(printed) < max(3, 0.1 * n):
        return pdf_labels, "pdf" if any(l != str(k + 1) for k, l in enumerate(pdf_labels)) else "physical"
    agree = sum(pdf_labels[i] == str(v) for i, v in printed.items())
    if agree >= 0.5 * len(printed):
        return pdf_labels, "pdf"
    out = list(pdf_labels)
    known = sorted(printed)
    import bisect
    for i in range(n):
        if i in printed:
            out[i] = str(printed[i])
            continue
        j = bisect.bisect_left(known, i)
        nb = [k for k in (known[j - 1] if j > 0 else None, known[j] if j < len(known) else None) if k is not None and abs(k - i) <= 6]
        for k in nb:
            if k != nb[0] and printed[k] - k != printed[nb[0]] - nb[0]:
                nb = []
                break
        if nb:
            out[i] = str(printed[nb[0]] + (i - nb[0]))
    return out, "printed"


def extract_pdf(path: Path,progress: Callable[[int, int], None] | None = None) -> dict:
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    pdf = pdfium.PdfDocument(str(path))
    try:
        n = len(pdf)
        labels: list[str] = []
        pages_lines: list[list[tuple[str, float, int, int]]] = []
        scanned: list[int] = []
        for i in range(n):
            lab = ""
            try:
                lab = pdf.get_page_label(i).strip()
            except Exception:  # noqa: BLE001
                pass
            labels.append(lab or str(i + 1))
            page = pdf[i]
            tp = page.get_textpage()
            txt = _clean_text(tp.get_text_range())
            lines = _line_heights(tp, raw, txt)
            if sum(len(l[0].strip()) for l in lines) < 20:
                scanned.append(i)
            pages_lines.append(lines)
            tp.close()
            page.close()
            if progress:
                progress(i + 1, n)
        pages_lines = _repair_ocr_l(pages_lines)
        toc = []
        try:
            for bm in pdf.get_toc():
                dest = bm.get_dest()
                idx = dest.get_index() if dest is not None else None
                title = _norm_space(_clean_text(bm.get_title() or "")).rstrip("*").strip()
                if idx is not None and title:
                    toc.append((bm.level, title, idx))
        except Exception as e:  # noqa: BLE001
            log.warning("outline read failed for %s: %s", path.name, e)
    finally:
        pdf.close()

    # repeated header/footer lines (first/last 2 non-empty lines; digits normalised)
    def key(s: str) -> str:
        return re.sub(r"\d+", "#", s.strip().lower())

    cnt: Counter = Counter()
    for lines in pages_lines:
        ne = [l[0] for l in lines if l[0].strip()]
        for s in set(ne[:2] + ne[-2:]):
            cnt[key(s)] += 1
    rep = {k for k, c in cnt.items() if n >= 4 and c >= max(3, 0.3 * n) and len(k) < 120}
    labels, label_src = _printed_labels(pages_lines, labels)

    heights = [l[1] for lines in pages_lines for l in lines if len(l[0].split()) >= 6 and l[1] > 0]
    body = statistics.median(heights) if heights else 0.0

    use_toc = len(toc) >= 2
    toc_by_page: dict[int, list[tuple[int, str]]] = defaultdict(list)
    for lvl, title, idx in toc:
        toc_by_page[idx].append((lvl + 1, title))
    # heading levels for the font heuristic: distinct big sizes, biggest = level 1
    big = sorted({round(l[1], 0) for lines in pages_lines for l in lines
                  if body and l[1] >= 1.18 * body and 0 < len(l[0].split()) <= 14}, reverse=True)
    lvl_of = {h: min(i + 1, 4) for i, h in enumerate(big)}

    def normt(s: str) -> str:
        return re.sub(r"[^a-z0-9]", "", s.lower())

    # running heads ("540 CHAPTER TITLE" / "SECTION TITLE 521") as a section fallback when the
    # PDF has neither an outline nor usable font-size headings (typical for OCR'd scans)
    run_heads: dict[int, str] = {}
    n_font = 0
    for i, lines in enumerate(pages_lines):
        ne = [l for l in lines if l[0].strip()]
        lab = labels[i]
        for l in ne[:2] + ne[-2:]:
            st = l[0].strip()
            if lab.isdigit() and len(st.split()) <= 12 and re.match(rf"^{lab}\b|.*\b{lab}$", st):
                t = _norm_space(re.sub(rf"^{lab}\b|\b{lab}$", "", st)).strip(" .|-")
                t = re.sub(r"[^A-Za-z ,'-]", " ", t)  # OCR margin debris
                t = _norm_space(" ".join(x for x in t.split() if not re.search(r"[a-z][A-Z]", x) and (len(x) > 1 or x in "aA"))).strip(" ,'-")
                if _heading_like(t) and len(t.split()) <= 10:
                    run_heads[i] = t.title() if t.isupper() else t
                    break
        if not use_toc:
            n_font += sum(1 for l in ne if body and l[1] >= 1.18 * body and _heading_like(l[0].strip()) and len(l[0].split()) <= 14)
    head_n = Counter(run_heads.values())
    run_heads = {k: v for k, v in run_heads.items() if head_n[v] >= 2 or k % 2 == (0 if sum(1 for x in run_heads if x % 2 == 0) else 1) and False}
    use_run = not use_toc and n_font < max(5, n / 30) and len(run_heads) >= 0.3 * n
    # part names ("Osteology") recur on many pages; section heads ("The Femur") on a few
    head_n = Counter(run_heads.values())
    parent_cut = max(12, sorted(head_n.values())[int(0.9 * (len(head_n) - 1))] if head_n else 12)
    cur_parent, cur_spec = None, None

    # boilerplate phrases that sit in every page's footer (a repeated page footer) are not evidence of a credits page
    page_hits = [{m.group(0).lower() for m in _BOILER_RX.finditer(" ".join(l[0] for l in lines))} for lines in pages_lines]
    common = {h for h, c in Counter(h for ph in page_hits for h in ph).items() if c > 0.4 * n and n > 3}
    blocks: list[tuple] = []
    if use_run and n >= 100:  # a book's pages before its first running head are title, preface, contents: stored, not indexed
        blocks.append(("h", 8, "Front matter", 0))
    dropped: dict[str, int] = {"toc_index": 0}
    for i, lines in enumerate(pages_lines):
        if use_run and i in run_heads:
            if head_n[run_heads[i]] < parent_cut:
                if run_heads[i] != cur_spec:
                    cur_spec = run_heads[i]
                    blocks.append(("h", _special(cur_spec, 2), cur_spec, i))
            elif run_heads[i] != cur_parent:
                cur_parent, cur_spec = run_heads[i], None
                blocks.append(("h", _special(cur_parent, 1), cur_parent, i))
        ne_lines = [l[0].strip() for l in lines if l[0].strip()]
        boiler = len(page_hits[i] - common)
        near_edge = n <= 40 and (i == 1 or i >= n - 3)  # short booklets: credits up front, blurbs/bios at the back
        if n > 3 and (boiler >= 2 or (boiler >= 1 and near_edge)):
            dropped["boilerplate"] = dropped.get("boilerplate", 0) + 1  # credits, licence blurbs, publisher back matter
            for lvl, title in (toc_by_page.get(i, []) if use_toc else []):
                blocks.append(("h", _special(title, lvl), title, i))
            continue
        if i > 0 and len(ne_lines) >= 12 and sum(bool(re.search(r"(\.{3,}|\s)\d{1,4}\s*$|,\s*\d{1,4}\s*$", x)) for x in ne_lines) >= 0.5 * len(ne_lines):
            dropped["toc_index"] += 1  # table of contents or back-of-book index: page references, not content
            for lvl, title in (toc_by_page.get(i, []) if use_toc else []):
                blocks.append(("h", _special(title, lvl), title, i))
            continue
        pending = list(toc_by_page.get(i, [])) if use_toc else []
        nonempty = [k for k, l in enumerate(lines) if l[0].strip()]
        edge = set(nonempty[:2] + nonempty[-2:])
        placed = set()
        for k, (s, h, w, _off) in enumerate(lines):
            st = s.strip()
            if not st:
                blocks.append(("t", 0, "", i))  # paragraph break marker
                continue
            if k in edge and (key(st) in rep or _PAGENUM_RX.match(st)):
                continue
            if k in edge and len(st.split()) <= 12 and labels[i].isdigit() and re.match(rf"^{labels[i]}\b|.*\b{labels[i]}$", st):
                continue  # running head carrying the printed page number ("212 CHAPTER TITLE")
            if use_toc:
                ns = normt(st)
                hit = None
                for j, (lvl, title) in enumerate(pending):
                    nt = normt(title)
                    if len(ns) >= 4 and (nt == ns or (nt.startswith(ns) and len(ns) >= 8) or (ns.startswith(nt) and len(nt) >= 6 and len(st.split()) <= 16)):
                        hit = j
                        break
                if hit is not None:
                    lvl, title = pending.pop(hit)
                    blocks.append(("h", _special(title, lvl), title, i))
                    placed.add(title)
                    continue
                if _special(st, 0):
                    blocks.append(("h", _special(st, 0), st, i))
                    continue
            else:
                words = st.split()
                if _special(st, 0):
                    blocks.append(("h", _special(st, 0), st, i))
                    continue
                is_head = (not use_run and body and h >= 1.18 * body and _heading_like(st) and len(words) <= 14 and re.search(r"[A-Za-z]{2}", st)
                           and not st.endswith((".", ",", ";")) and not _PAGENUM_RX.match(st))
                if is_head:
                    blocks.append(("h", lvl_of.get(round(h, 0), 3), st, i))
                    continue
                if (not use_run and w >= 650 and len(words) <= 10 and _heading_like(st) and not st.endswith((".", ",", ";", ":")) and re.match(r"^(\d+(\.\d+)*\s+)?[A-Z]", st)
                        and (k + 1 >= len(lines) or not lines[k + 1][0].strip() or lines[k + 1][2] < 650)):
                    blocks.append(("h", 5, st, i))
                    continue
                if _special(st, 0):
                    blocks.append(("h", _special(st, 0), st, i))
                    continue
            bm = _BACKMATTER_RX.match(st) if len(st) < 400 else None
            if bm and bm.group(3) and st[:1].isupper() and (not bm.group("rest") or bm.group("rest")[:1].isupper()):
                blocks.append(("h", 9, st[: len(st) - len(bm.group("rest") or "")].strip(), i))
                if bm.group("rest"):
                    blocks.append(("t", 0, bm.group("rest"), i))
                continue
            blocks.append(("t", 0, st, i))
        # outline entries whose line was not found: start the section at the top of the page
        pos = next((b for b in range(len(blocks)) if blocks[b][3] == i), len(blocks))
        for lvl, title in pending:
            blocks.insert(pos, ("h", _special(title, lvl), title, i))
            pos += 1
    return {"blocks": blocks, "labels": labels, "pages": n, "scanned": scanned,
            "page_labels": label_src, "dropped_pages": dropped, "headings": "outline" if use_toc else ("running_heads" if use_run else ("font" if big else "none"))}


def _strip_md(s: str) -> str:
    s = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"`{1,3}|\*\*|__|(?<!\w)[*_](?=\w)|(?<=\w)[*_](?!\w)", "", s)
    s = re.sub(r"^\s*(>|[-*+]\s|\d+\.\s)", "", s)
    s = re.sub(r"<[^>]+>", "", s)
    return s


def extract_text(path: Path, markdown: bool) -> dict:
    raw_text = _clean_text(Path(path).read_text(errors="replace"))
    blocks: list[tuple] = []
    page = 0
    labels: list[str] = []
    page_rx = re.compile(r"^\s*(<!--\s*page[:\s]*([\w.-]+)\s*-->|\[\s*p(?:age)?\.?\s*([\w.-]+)\s*\])\s*$", re.I)
    lines = raw_text.split("\n")
    in_code = False
    table_hdr: list[str] | None = None
    for k, line in enumerate(lines):
        if "\f" in line:
            page += 1
            line = line.replace("\f", "")
        m = page_rx.match(line)
        if m:
            labels.append(m.group(2) or m.group(3))
            page = len(labels) - 1
            continue
        st = line.strip()
        if markdown and st.startswith("```"):
            in_code = not in_code
            continue
        if not st:
            blocks.append(("t", 0, "", page))
            continue
        if markdown and st.startswith("|") and not in_code:
            # Markdown table → one self-describing line per row ("Day: Mon 12 Oct · Subject: Science · Homework: …")
            cells = [_strip_md(c.strip()) for c in st.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells if c):
                continue
            nxt = lines[k + 1].strip() if k + 1 < len(lines) else ""
            if table_hdr is None and re.match(r"^\|?\s*:?-{2,}", nxt):
                table_hdr = cells
                continue
            if table_hdr:
                row = " · ".join(f"{h}: {c}" if h else c for h, c in zip(table_hdr + [""] * len(cells), cells) if c)
            else:
                row = " · ".join(c for c in cells if c)
            blocks.append(("t", 0, row.rstrip(".") + ".", page))
            blocks.append(("t", 0, "", page))
            continue
        table_hdr = None if not st.startswith("|") else table_hdr
        if markdown:
            hm = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", st)
            if hm and not in_code:
                lvl = len(hm.group(1))
                t = _strip_md(hm.group(2))
                blocks.append(("h", _special(t, lvl), t, page))
                continue
            if re.fullmatch(r"[-=*_|:\s]{3,}", st):
                continue
            st = _strip_md(st)
        else:
            prev_blank = k == 0 or not lines[k - 1].strip()
            next_blank = k + 1 >= len(lines) or not lines[k + 1].strip()
            if prev_blank and next_blank and len(st.split()) <= 10 and not st.endswith((".", ",", ";", ":", "?", "!")) and re.search(r"[A-Za-z]", st):
                blocks.append(("h", _special(st, 1 if st.isupper() else 2), st, page))
                continue
        blocks.append(("t", 0, st, page))
    npages = max(page + 1, len(labels), 1)
    has_pages = bool(labels) or page > 0
    if not labels:
        labels = [str(i + 1) for i in range(npages)] if has_pages else [""]
    return {"blocks": blocks, "labels": labels, "pages": npages, "scanned": [],
            "headings": "markdown" if markdown else "text"}


# ───────────────────────────── spreadsheets ─────────────────────────────

def _parse_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%d-%b-%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
            try:
                return dt.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


# Currency columns: the symbol is shown only when the column header names the currency.
_CURRENCIES = (("₹", r"inr|rs\.?|rupees?|₹"), ("$", r"usd|dollars?|\$"), ("€", r"eur|euros?|€"), ("£", r"gbp|pounds?|£"))
_CUR_STRIP_RX = re.compile(r"[₹$€£]|(?<![a-z])(?:rs\.|inr|usd|eur|gbp)(?![a-z])", re.I)


def _currency_symbol(header: str) -> str:
    h = (header or "").lower()
    for sym, rx in _CURRENCIES:
        if re.search(r"(?<![a-z])(?:" + rx + r")(?![a-z])", h):
            return sym
    return ""


def _parse_num(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = _CUR_STRIP_RX.sub("", v.replace(",", "")).strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            return float(s)
    return None


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        d = v.date() if isinstance(v, dt.datetime) else v
        return f"{d.day} {_MONTHS[d.month - 1][:3]} {d.year}"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.2f}".rstrip("0").rstrip(".")
    return _norm_space(str(v))


def _fmt_num(x: float) -> str:
    return f"{int(round(x)):,}" if abs(x - round(x)) < 1e-6 else f"{x:,.2f}"


def read_sheets(path: Path) -> dict[str, list[list]]:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        finally:
            wb.close()
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return {path.stem: [list(r) for r in csv.reader(f, dialect)]}


def _cell(v):
    """Normalise a cell: ISO/other date strings become dates, so rows read '1 Jul 2026' whatever the source."""
    if isinstance(v, str):
        d = _parse_date(v) if re.match(r"^\d{4}-\d{2}-\d{2}|^\d{1,2}[/-]\d{1,2}[/-]\d{4}$", v.strip()) else None
        return d if d else v.strip()
    return v


ROW_CHARS = 360   # ≈ 70–90 MiniLM wordpieces for a compact spreadsheet row


def _cell_text(h: str, v) -> str:
    hl = h.lower()
    x = _parse_num(v) if not isinstance(v, (dt.date, dt.datetime)) else None
    sym = _currency_symbol(hl) if x is not None else ""
    if sym:
        label = re.sub(r"\s*\(?\s*(?:" + "|".join(rx for _, rx in _CURRENCIES) + r")\s*\)?", "", h, flags=re.I).strip()
        return f"{label} {sym}{_fmt_num(x)}"
    return f"{h}: {_fmt(v)}"


def _num_priority(h: str) -> int:
    """Which numeric columns to total first; 9 = never sum (ranks, rates, unit prices, ages)."""
    h = h.lower()
    if re.search(r"\b(rank|rate|pct|percent|level|price|mrp|days|lead|year|age|margin|score|no|number)\b|%", h):
        return 9
    if re.search(r"\b(total|amount|refund|revenue|sales|value)\b", h):
        return 0
    if re.search(r"\b(qty|quantity|units|stock|count|orders)\b", h):
        return 1
    return 2


def sheet_chunks(sheet: str, rows: list[list]) -> list[dict]:
    """Row chunks + summary chunks for one sheet. Each: {text, label, kind, row}.

    Rows: "Orders row 42: order id: ORD-1042 · order date: 1 Jul 2026 · channel: web · …".
    Summaries (so totals never need the LLM to add up rows): totals per sheet; counts and sums by each
    categorical column; by month of each date column (with the IDs when ≤ 25); per product type derived from
    the shared tail of item names ("… Rain Jacket") with every numeric column summed (stock by size, etc.)."""
    hdr_i = None
    for i, r in enumerate(rows[:10]):
        vals = [c for c in r if c not in (None, "")]
        if len(vals) >= 2 and sum(isinstance(c, str) for c in vals) >= 0.8 * len(vals):
            hdr_i = i
            break
    if hdr_i is None:  # a notes sheet ("Read_Me"): keep its text as one chunk
        lines = [" ".join(_fmt(c) for c in r if c not in (None, "")) for r in rows]
        text = " ".join(x for x in lines if x)
        return [{"text": f"{sheet}: {text}", "label": sheet, "kind": "note", "row": None}] if text else []
    hdr = [(_norm_space(str(c).replace("_", " ")) if c not in (None, "") else f"col{j + 1}") for j, c in enumerate(rows[hdr_i])]
    data = []
    for i, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):  # 1-based spreadsheet row numbers
        if any(c not in (None, "") for c in r):
            r = [_cell(c) for c in r] + [None] * (len(hdr) - len(r))
            data.append((i, r[:len(hdr)]))
    out: list[dict] = []
    n = len(data)
    cols: dict[int, str] = {}
    ndist: dict[int, int] = {}
    for j, h in enumerate(hdr):
        vals = [r[j] for _, r in data if r[j] not in (None, "")]
        if not vals:
            continue
        nums = [_parse_num(v) for v in vals]
        distinct = {_fmt(v) for v in vals}
        ndist[j] = len(distinct)
        if sum(_parse_date(v) is not None for v in vals) >= 0.8 * len(vals):
            cols[j] = "date"
        elif sum(x is not None for x in nums) >= 0.8 * len(vals):
            id_like = len(distinct) == len(vals) and re.search(r"\b(id|no|number|code|pin|phone|tracking|mobile)\b", h.lower())
            cols[j] = "id" if id_like else "num"
        elif len(distinct) == len(vals) and len(vals) >= 0.9 * n:
            cols[j] = "name" if re.search(r"product|item|title|description|style|design", h.lower()) else "id"
        elif 2 <= len(distinct) <= min(40, max(2, int(0.6 * n))):
            cols[j] = "cat"
        elif re.search(r"product|item|title|description|style|design", h.lower()):
            cols[j] = "name"
        else:
            cols[j] = "text"
    # ---- one compact row chunk per row (≈ 40–80 tokens): key columns first, tax/rate detail dropped when long
    def col_score(j: int) -> int:
        h, t = hdr[j].lower(), cols.get(j)
        if t is None:
            return 9
        if re.search(r"gst|tax|hsn|discount|\brate\b|pct|percent|mrp|unit price", h):
            return 5
        if j == first_id or t == "name" or re.search(r"\b(status|product|item)\b", h):
            return 0
        if t == "date" or (t == "num" and _num_priority(hdr[j]) <= 1) or re.search(r"tracking|courier", h):
            return 1
        if t == "cat":
            return 2 if ndist.get(j, 0) > 12 or re.search(r"payment|mode|colou?r", h) else 1
        if t in ("id", "text"):
            return 2
        return 3
    first_id = next((j for j, t in cols.items() if t == "id"), None)
    order = sorted(cols, key=lambda j: (col_score(j), j))
    for rn, r in data:
        keep, used = [], 0
        for j in order:
            v = r[j]
            if v in (None, "") or not _fmt(v) or (cols[j] in ("num",) and _parse_num(v) == 0 and col_score(j) >= 2):
                continue
            piece = _cell_text(hdr[j], v)
            if col_score(j) >= 5 or (used + len(piece) > ROW_CHARS and col_score(j) > 0):
                continue
            keep.append((j, piece))
            used += len(piece) + 3
        text = f"{sheet} row {rn}: " + " · ".join(pc for _, pc in sorted(keep))
        out.append({"text": text, "label": f"{sheet} row {rn}", "kind": "row", "row": rn})
    if n < 3:
        return out
    cats = sorted((j for j, t in cols.items() if t == "cat"), key=lambda j: ndist[j])
    nums = sorted((j for j, t in cols.items() if t == "num" and _num_priority(hdr[j]) < 9), key=lambda j: (_num_priority(hdr[j]), j))
    dates = [j for j, t in cols.items() if t == "date"]
    ids = [j for j, t in cols.items() if t == "id"]
    idcol = ids[0] if ids else None
    lab = f"{sheet} summary"
    allrows = [r for _, r in data]
    noun = _row_noun(sheet)  # "5 returns", not "5 rows": a small model reads "rows" next to refund totals as ambiguous

    def numsum(sel: list[list], j: int) -> float:
        return sum((_parse_num(r[j]) or 0.0) for r in sel)

    def agg_line(sel: list[list], k: int = 3) -> str:
        bits = [f"{len(sel)} {noun}"]
        for j in nums[:k]:
            bits.append(f"total {hdr[j]} {_fmt_num(numsum(sel, j))}")
        return ", ".join(bits)

    out.append({"text": f"{sheet} sheet summary: {n} rows. Columns: {', '.join(hdr)}. " +
                "; ".join(f"total {hdr[j]} {_fmt_num(numsum(allrows, j))}" for j in nums[:8]) + ".",
                "label": lab, "kind": "summary", "row": None})
    for j in cats:
        groups: dict[str, list] = defaultdict(list)
        for r in allrows:
            if r[j] not in (None, ""):
                groups[_fmt(r[j])].append(r)
        items = sorted(groups.items(), key=lambda kv: -len(kv[1]))
        body = "; ".join(f"{k}: {agg_line(v, 2)}" for k, v in items[:40])
        # say the winner in words first: a 2B model does not reliably pick the largest count out of a list
        lead = ""
        if items and (len(items) == 1 or len(items[0][1]) > len(items[1][1])):
            lead = f"Most {noun} by {hdr[j]}: {items[0][0]} ({len(items[0][1])} {noun}). "
        elif items:
            tied = [k for k, v in items if len(v) == len(items[0][1])]
            lead = f"Most {noun} by {hdr[j]}: a tie between {', '.join(tied[:4])} ({len(items[0][1])} {noun} each). "
        out.append({"text": f"{sheet} by {hdr[j]} ({len(items)} values, {n} {noun}): {lead}{body}.", "label": lab, "kind": "summary", "row": None})
    for j in dates:
        months: dict[tuple, list] = defaultdict(list)
        for r in allrows:
            d = _parse_date(r[j])
            if d:
                months[(d.year, d.month)].append(r)
        ms = sorted(months)
        body = "; ".join(f"{_MONTHS[m - 1]} {y}: {agg_line(months[(y, m)], 2)}" for y, m in ms)
        out.append({"text": f"{sheet} by month of {hdr[j]}: {body}.", "label": lab, "kind": "summary", "row": None})
        if j != dates[0]:
            continue  # per-month detail only for the sheet's main date column
        for y, m in ms[-24:]:
            sel = months[(y, m)]
            t = f"{sheet} in {_MONTHS[m - 1]} {y} ({_MONTHS[m - 1][:3]} {y}, by {hdr[j]}): {agg_line(sel)}."
            for c in [c for c in cats if ndist[c] <= 12][:4]:
                cc = Counter(_fmt(r[c]) for r in sel if r[c] not in (None, ""))
                t += f" By {hdr[c]}: " + ", ".join(f"{k} {v}" for k, v in cc.most_common(12)) + "."
            if idcol is not None and len(sel) <= 25:
                t += f" {hdr[idcol]}s: " + ", ".join(_fmt(r[idcol]) for r in sel) + "."
            out.append({"text": t, "label": lab, "kind": "summary", "row": None})
    # product types from the shared tail of item names: "Navy Hooded Rain Jacket" → "Rain Jacket"
    for j in [j for j, t in cols.items() if t == "name"][:1]:
        names = [_fmt(r[j]) for r in allrows]
        tails2 = Counter(" ".join(x.split()[-2:]) for x in names if len(x.split()) >= 3)
        tails1 = Counter(x.split()[-1] for x in names if x.split())
        groups = defaultdict(list)
        for x, r in zip(names, allrows):
            w = x.split()
            t2 = " ".join(w[-2:]) if len(w) >= 3 else None
            key = t2 if t2 and tails2[t2] >= 2 and tails1[w[-1]] == tails2[t2] else (w[-1] if w and tails1[w[-1]] >= 2 else None)
            if key:
                groups[key].append(r)
        for key, sel in list(groups.items())[:60]:
            t = (f"{sheet}: all {key} items together ({len(sel)} {hdr[j]}s: " + ", ".join(_fmt(r[j]) for r in sel[:12]) + "). Totals: " +
                 "; ".join(f"{hdr[nj]} {_fmt_num(numsum(sel, nj))}" for nj in sorted(nums)[:16]) + ".")
            out.append({"text": t, "label": lab, "kind": "summary", "row": None})
    # long-format pivots: entity column × small categorical (e.g. product × size) → numeric sums
    small = [j for j in cats if ndist[j] <= 8]
    budget = 150
    for a in [j for j in cats if j not in small][:2]:
        ga: dict[str, list] = defaultdict(list)
        for r in allrows:
            if r[a] not in (None, ""):
                ga[_fmt(r[a])].append(r)
        for av, sel in ga.items():
            if budget <= 0 or not small:
                break
            t = f"{sheet}: {hdr[a]} {av} — {agg_line(sel)}."
            for b in small[:2]:
                gb: dict[str, list] = defaultdict(list)
                for r in sel:
                    gb[_fmt(r[b])].append(r)
                if nums:
                    t += f" {hdr[nums[0]]} by {hdr[b]}: " + ", ".join(f"{k} {_fmt_num(numsum(v, nums[0]))}" for k, v in gb.items()) + "."
                else:
                    t += f" Rows by {hdr[b]}: " + ", ".join(f"{k} {len(v)}" for k, v in gb.items()) + "."
            out.append({"text": t, "label": lab, "kind": "summary", "row": None})
            budget -= 1
    out.extend(_explicit_summaries(sheet, hdr, cols, ndist, data, nums, dates, cats, lab))
    return out


def _row_noun(sheet: str) -> str:
    t = sheet.lower()
    for key, noun in (("return", "returns"), ("order", "orders"), ("inventor", "products"), ("stock", "products"),
                      ("customer", "customers"), ("sale", "sales"), ("supplier", "suppliers")):
        if key in t:
            return noun
    return "rows"


def _explicit_summaries(sheet, hdr, cols, ndist, data, nums, dates, cats, lab) -> list[dict]:
    """Plain sentences for the questions people actually ask: best sellers, top customer, items below reorder, by quarter."""
    out: list[dict] = []
    allrows = [r for _, r in data]
    ids = [j for j, t in cols.items() if t == "id"]
    names = [j for j, t in cols.items() if t in ("name", "cat") and re.search(r"product|item|style|design|description", hdr[j].lower())]
    qty = next((j for j in nums if re.search(r"\b(qty|quantity|units)\b", hdr[j].lower())), None)
    money = next((j for j in nums if re.search(r"\btotal\b", hdr[j].lower()) and not re.search(r"stock|qty|units", hdr[j].lower())),
                 next((j for j in nums if re.search(r"amount|revenue|sales|value|refund", hdr[j].lower())), None))
    today = dt.date.today()

    def numsum(sel, j):
        return sum((_parse_num(r[j]) or 0.0) for r in sel)

    def label_row(r) -> str:
        bits = [_fmt(r[j]) for j in (ids[:1] + names[:1]) if r[j] not in (None, "")]
        return " ".join(bits) or "row"

    def top_sellers(sel, scope: str) -> str | None:
        if not names or qty is None:  # needs a units-sold column (an inventory sheet has none)
            return None
        g = defaultdict(list)
        for r in sel:
            if r[names[0]] not in (None, ""):
                g[_fmt(r[names[0]])].append(r)
        key = (lambda kv: (numsum(kv[1], qty), numsum(kv[1], money) if money is not None else 0)) if qty is not None else (lambda kv: numsum(kv[1], money))
        best = sorted(g.items(), key=key, reverse=True)[:8]
        items = []
        for i, (nm, rs) in enumerate(best, 1):
            bits = []
            if qty is not None:
                bits.append(f"{_fmt_num(numsum(rs, qty))} units")
            if money is not None:
                cur = _currency_symbol(hdr[money])
                bits.append(f"{cur}{_fmt_num(numsum(rs, money))}")
            items.append(f"{i}. {nm} — {', '.join(bits)}")
        return f"Top sellers (best-selling products{scope}, from {sheet}): " + "; ".join(items) + "."

    t = top_sellers(allrows, ", all rows")
    if t:
        out.append({"text": t, "label": lab, "kind": "summary", "row": None})
    if dates:
        qs = defaultdict(list)
        for r in allrows:
            d = _parse_date(r[dates[0]])
            if d:
                qs[(d.year, (d.month - 1) // 3 + 1)].append(r)
        cur_q = (today.year, (today.month - 1) // 3 + 1)
        done = [q for q in sorted(qs) if q < cur_q]
        for (y, qn) in sorted(qs):
            sel = qs[(y, qn)]
            m0 = (qn - 1) * 3
            span = f"{_MONTHS[m0]}–{_MONTHS[m0 + 2]} {y}"
            rel = " (last quarter)" if done and (y, qn) == done[-1] else (" (this quarter, so far)" if (y, qn) == cur_q else "")
            body = f"{sheet} in Q{qn} {y}{rel}, {span}: {len(sel)} {_row_noun(sheet)}"
            if qty is not None:
                body += f", {_fmt_num(numsum(sel, qty))} units"
            if money is not None:
                body += f", total {hdr[money]} {_currency_symbol(hdr[money])}{_fmt_num(numsum(sel, money))}"
            body += "."
            ts = top_sellers(sel, f" in Q{qn} {y}{rel}")
            if ts:
                body += " " + ts
            out.append({"text": body, "label": lab, "kind": "summary", "row": None})
    # ranking sheets ("Top customers": rank column)
    rank = next((j for j, h in enumerate(hdr) if re.fullmatch(r"rank(ing)?|#|no\.?", h.lower().strip())), None)
    if rank is not None:
        ranked = sorted((r for r in allrows if _parse_num(r[rank]) is not None), key=lambda r: _parse_num(r[rank]))[:5]
        noun = re.sub(r"[_\s]*top\s*\d+|[_\s]+\d+$", "", sheet, flags=re.I).replace("_", " ").strip().lower()
        noun = noun[:-1] if noun.endswith("s") else noun
        def rtext(r):
            return " · ".join(_cell_text(hdr[j], r[j]) for j in cols if j != rank and r[j] not in (None, ""))
        if ranked:  # rank 1 on its own (short chunk embeds close to "who's my top customer?"), then the runners-up
            out.append({"text": f"Top {noun} (best {noun}, number one, rank 1 in {sheet}): {rtext(ranked[0])}.", "label": lab, "kind": "summary", "row": None})
        if len(ranked) > 1:
            out.append({"text": f"Top {len(ranked)} {noun}s in {sheet}, after number one: " + "; ".join(f"{i}. {rtext(r)}" for i, r in enumerate(ranked[1:], 2)) + ".",
                        "label": lab, "kind": "summary", "row": None})
    # yes/no flag columns ("below reorder": yes) → the list of flagged items
    for j, tp in cols.items():
        vals = {str(r[j]).strip().lower() for r in allrows if r[j] not in (None, "")}
        if tp == "cat" and vals <= {"yes", "no", "y", "n", "true", "false"} and len(vals) >= 1:
            hit = [r for r in allrows if str(r[j]).strip().lower() in ("yes", "y", "true")]
            extra = [k for k in nums[:2]]
            def ftext(r):
                return label_row(r) + (" (" + ", ".join(_cell_text(hdr[k], r[k]) for k in extra if r[k] not in (None, "")) + ")" if extra else "")
            t = (f"{sheet}: items where {hdr[j]} is yes ({len(hit)} of {len(allrows)}): " + ("; ".join(ftext(r) for r in hit[:20]) if hit else "none") + ".")
            out.append({"text": t, "label": lab, "kind": "summary", "row": None})
    # small categorical groups: name their rows ("status processing: ORD-1281, ORD-1283 …")
    if ids:
        for j in cats:
            if ndist[j] > 12:
                continue
            g = defaultdict(list)
            for r in allrows:
                if r[j] not in (None, ""):
                    g[_fmt(r[j])].append(r)
            small = [(k, v) for k, v in g.items() if len(v) <= 12]
            if small:
                t = f"{sheet} — rows by {hdr[j]}: " + "; ".join(f"{k} ({len(v)}): " + ", ".join(label_row(r) for r in v) for k, v in small) + "."
                out.append({"text": t, "label": lab, "kind": "summary", "row": None})
    return out


# ───────────────────────────── chunking ─────────────────────────────

class _Tok:
    def __init__(self) -> None:
        self.tok = embedder.get().tok

    def counts(self, texts: list[str]) -> list[int]:
        if not texts:
            return []
        return [len(e.ids) - 2 for e in self.tok.encode_batch(texts)]


def build_chunks(ext: dict, short: str, tok: _Tok) -> list[dict]:
    """Blocks → section-aware, sentence-aligned chunks. Each chunk: kind, section path, sentences[(text, page)]."""
    labels = ext["labels"]
    path: list[tuple[int, str]] = []
    sections: list[dict] = []
    cur = {"path": [], "refs": False, "paras": []}
    para: list[tuple[str, int]] = []

    def flush_para():
        nonlocal para
        if para:
            cur["paras"].append(para)
            para = []

    for kind, lvl, text, page in ext["blocks"]:
        if kind == "h":
            flush_para()
            if cur["paras"]:
                sections.append(cur)
            if lvl >= 8:
                cur = {"path": [t for _, t in path] + [text], "refs": True, "paras": []}
                continue
            while path and path[-1][0] >= lvl:
                path.pop()
            path.append((lvl, text))
            cur = {"path": [t for _, t in path], "refs": False, "paras": []}
            continue
        if not text:
            flush_para()
            continue
        if para and re.search(r"[A-Za-z]-$", para[-1][0]) and re.match(r"[a-z]", text):
            # de-hyphenate "physio-" + "logy is ..." → "physiology is ..."
            p0, pg = para[-1]
            first, _, rest = text.partition(" ")
            para[-1] = (p0[:-1] + first, pg)
            if rest:
                para.append((rest, page))
        else:
            para.append((text, page))
    flush_para()
    if cur["paras"]:
        sections.append(cur)

    chunks: list[dict] = []
    for sec in sections:
        sents: list[tuple[str, int]] = []
        for para in sec["paras"]:
            # join lines, remembering the page where each character run starts
            buf, marks = "", []
            for t, pg in para:
                marks.append((len(buf), pg))
                buf += (" " if buf else "") + t
            for s in split_sentences(buf):
                pos = buf.find(s[:20])
                pg = marks[0][1]
                for o, p in marks:
                    if o <= max(pos, 0):
                        pg = p
                sents.append((s, pg))
        sents = [s for s in sents if re.search(r"[A-Za-z]{2}", s[0])]
        if not sents:
            continue
        ntoks = tok.counts([s for s, _ in sents])
        groups: list[list[int]] = [[]]
        acc = 0
        for i, nt in enumerate(ntoks):
            if groups[-1] and acc + nt > CHUNK_MAX_TOK:
                groups.append([])
                acc = 0
            groups[-1].append(i)
            acc += nt
        if len(groups) > 1 and sum(ntoks[i] for i in groups[-1]) < CHUNK_MIN_TOK and sum(ntoks[i] for i in groups[-2]) + sum(ntoks[i] for i in groups[-1]) <= CHUNK_MAX_TOK + 60:
            groups[-2].extend(groups.pop())
        section = " › ".join(sec["path"])
        for g in groups:
            ss = [(sents[i][0], sents[i][1], ntoks[i]) for i in g]
            pg = ss[0][1]
            chunks.append({
                "kind": "refs" if sec["refs"] else "prose",
                "section": section,
                "header": f"{short} › {section}" if section else short,
                "sents": [(s, p, labels[p] if p < len(labels) else str(p + 1), n) for s, p, n in ss],
                "page": pg, "page_label": labels[pg] if pg < len(labels) else str(pg + 1),
                "page_end": ss[-1][1],
            })
    return chunks


# ───────────────────────────── spoken-ID query normaliser ─────────────────────────────
# Speech recognition renders IDs like "ORD-1042" as "o r d ten forty two" etc.
# Prefixes, widths, number ranges and SKU categories are learned from the library's own IDs at load time,
# so the same rule works for any library. Only the search query is rewritten.

_UNITS = {"zero": 0, "oh": 0, "o": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
          "seven": 7, "eight": 8, "nine": 9}
_TEENS = {"ten": 10, "eleven": 11, "twelve": 12, "thirteen": 13, "fourteen": 14, "fifteen": 15,
          "sixteen": 16, "seventeen": 17, "eighteen": 18, "nineteen": 19}
_TENS = {"twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60, "seventy": 70, "eighty": 80, "ninety": 90}
_NUMWORDS = set(_UNITS) | set(_TEENS) | set(_TENS) | {"hundred", "thousand", "double", "triple"}
_LETTER = {"and": "m", "em": "m", "emm": "m", "en": "n", "an": "n", "kay": "k", "jay": "j", "dee": "d", "bee": "b",
           "see": "c", "pee": "p", "tee": "t", "vee": "v", "ess": "s", "es": "s", "ex": "x", "why": "y", "are": "r",
           "el": "l", "ef": "f", "aitch": "h", "gee": "g", "zed": "z", "zee": "z", "you": "u", "cue": "q"}
_ORDER_WORDS = {"order", "orders", "order's", "id", "number"}
_ID_TOK = re.compile(r"[A-Z]{2,5}-[A-Z]{2}\d{2}\b|[A-Za-z]+(?:'[a-z]+)?|\d+|[^\sA-Za-z\d]")


def _id_tokens(text: str) -> list[str]:
    text = re.sub(r"(?<![-\w])([A-Za-z]{1,5})-?(\d{2,6})\b", r"\1 \2", text)
    return _ID_TOK.findall(text)


def _parse_number(w: list[str]) -> str | None:
    if all(x.isdigit() for x in w):
        return "".join(w)
    out: list[str] = []
    k = 0
    if "thousand" in w or "hundred" in w:
        total, cur = 0, 0
        for x in w:
            if x.isdigit():
                cur += int(x)
            elif x in _UNITS:
                cur += _UNITS[x]
            elif x in _TEENS:
                cur += _TEENS[x]
            elif x in _TENS:
                cur += _TENS[x]
            elif x == "hundred":
                cur = max(cur, 1) * 100
            elif x == "thousand":
                total += max(cur, 1) * 1000
                cur = 0
        return str(total + cur)
    if all(x in _UNITS or x in ("double", "triple") or (x.isdigit() and len(x) == 1) for x in w):
        while k < len(w):
            x = w[k]
            if x in ("double", "triple") and k + 1 < len(w):
                out.append(str(_UNITS.get(w[k + 1], w[k + 1])) * (2 if x == "double" else 3))
                k += 2
                continue
            out.append(str(_UNITS.get(x, x)))
            k += 1
        return "".join(out)
    while k < len(w):
        x = w[k]
        if x.isdigit():
            out.append(x)
        elif x in _TEENS:
            out.append(str(_TEENS[x]))
        elif x in _TENS:
            v = _TENS[x]
            if k + 1 < len(w) and w[k + 1] in _UNITS and _UNITS[w[k + 1]] > 0:
                v += _UNITS[w[k + 1]]
                k += 1
            out.append(str(v))
        elif x in ("oh", "o", "zero") and k + 1 < len(w) and w[k + 1] in _UNITS:
            out.append("0" + str(_UNITS[w[k + 1]]))
            k += 1
        elif x in _UNITS:
            out.append(str(_UNITS[x]))
        k += 1
    return "".join(out) or None


def _num_run(toks: list[str], i: int) -> tuple[int, str | None]:
    j, words = i, []
    while j < len(toks):
        t = toks[j].lower()
        if t in (",", "-"):
            j += 1
            continue
        if t.isdigit() or t in _NUMWORDS:
            words.append(t)
            j += 1
            continue
        break
    while j > i and toks[j - 1] in (",", "-"):
        j -= 1
    return (j, _parse_number(words)) if words else (i, None)


def _letters(tok: str) -> str | None:
    t = tok.lower()
    if t in _LETTER:
        return _LETTER[t]
    return t if re.fullmatch(r"[a-z]{1,5}", t) else None


def normalize_ids(text: str, prefixes: dict, bare: tuple[str, tuple[int, int]] | None = None) -> str:
    """'Where is order o r d ten forty two' → 'Where is order ORD-1042' for prefixes {'ORD': (4, (1000, 1999))}."""
    if not prefixes:
        return text
    toks = _id_tokens(text)
    out: list[str] = []
    i = 0
    maxlen = max(len(p) for p in prefixes)
    while i < len(toks):
        hit = None
        for span in range(1, 9) if toks[i] not in (",", "-") else ():
            if i + span > len(toks) or toks[i + span - 1] in (",", "-"):
                continue
            seg = [t for t in toks[i:i + span] if t not in (",", "-")]
            ls = [_letters(t) for t in seg]
            if None in ls:
                break
            cand = "".join(ls).upper()
            if cand in prefixes:
                end, num = _num_run(toks, i + span)
                if num:
                    width, rng = prefixes[cand]
                    if rng is None or rng[0] <= int(num) <= rng[1]:
                        hit = (end, f"{cand}-{num.zfill(width)}")
                break
            if len(cand) >= maxlen:
                break
        if hit is None and bare and toks[i].lower() in _ORDER_WORDS:
            j = i + 1
            if j < len(toks) and toks[j].lower() in ("number", "no", "id"):
                j += 1
            end, num = _num_run(toks, j)
            if num and bare[1][0] <= int(num) <= bare[1][1]:
                out.extend(toks[i:j])
                hit = (end, f"{bare[0]}-{num}")
                i = j
        if hit:
            out.append(hit[1])
            i = hit[0]
        else:
            out.append(toks[i])
            i += 1
    s = " ".join(out)
    return re.sub(r"\s+([,.?!;:])", r"\1", s).replace(" ' ", "'").replace(" '", "'")


def normalize_skus(text: str, prefix: str, cats: list[str]) -> str:
    """'ordj k01' / 'O R D L S zero three' → 'ORD-JK01' / 'ORD-LS03'."""
    if not prefix or not cats:
        return text
    sep = r"[\s,._'-]*"
    sp = lambda w: sep.join(re.escape(ch) for ch in w.lower())  # noqa: E731
    rx = re.compile(r"\b(?P<pre>" + sp(prefix) + sep + r")?(?P<cat>" + "|".join(sp(c) for c in cats) + r")" + sep
                    + r"(?P<zero>(?:0|zero|oh|o)" + sep + r")?(?P<u>one|two|three|four|five|six|seven|eight|nine|[1-9])\b", re.I)

    def rep(m):
        if not (m.group("pre") or m.group("zero")):
            return m.group(0)
        cat = re.sub(r"[^a-z]", "", m.group("cat").lower()).upper()
        u = m.group("u")
        return f"{prefix}-{cat}0{_UNITS.get(u.lower(), u)}"
    return rx.sub(rep, text)


def learn_id_scheme(texts) -> dict:
    """From indexed text: ID prefixes with digit width and range ('ORD-1042' → ORD: 4 digits, 1000–1999) and SKU categories."""
    nums: dict[str, list[str]] = defaultdict(list)
    skus: dict[str, Counter] = defaultdict(Counter)
    for t in texts:
        for m in re.finditer(r"\b([A-Z]{2,5})-(\d{2,6})\b", t):
            nums[m.group(1)].append(m.group(2))
        for m in re.finditer(r"\b([A-Z]{2,5})-([A-Z]{2})(\d{2})\b", t):
            skus[m.group(1)][m.group(2)] += 1
    prefixes = {}
    for p, ds in nums.items():
        if len(set(ds)) >= 5:
            w = Counter(len(d) for d in ds).most_common(1)[0][0]
            vals = sorted(int(d) for d in ds if len(d) == w)
            lo, hi = vals[0], vals[-1]
            span = 10 ** (w - 1)
            prefixes[p] = (w, (lo // span * span, (hi // span + 1) * span - 1))
    bare = None
    if prefixes:
        top = max(prefixes, key=lambda p: len(set(nums[p])))
        if prefixes[top][0] == 4:
            bare = (top, prefixes[top][1])
    sku_prefix = max(skus, key=lambda p: sum(skus[p].values())) if skus else ""
    return {"prefixes": prefixes, "bare": bare, "sku_prefix": sku_prefix,
            "sku_cats": sorted(skus[sku_prefix]) if sku_prefix else []}


# How-to-answer wording ("explain in detail, step by step, with all the main parts") says nothing about WHAT
# to find; it dilutes the query embedding and the keyword cover, so it is removed before searching.
_STYLE_RX = re.compile(
    r"\b(?:(?:can|could|would) you |please )?(?:explain|describe)(?= )|\b(?:in|and) (?:great |more |full |a lot of )?detail\b|\bdetailed\b"
    r"|\bstep[- ]by[- ]step\b|\bwith (?:all )?(?:the )?(?:main|domain) (?:parts|points|steps)\b|\bin (?:simple|plain|easy) (?:words|terms|language)\b"
    r"|\b(?:briefly|quickly|simply|please)\b|\bin (?:one|a few|two|three) (?:lines?|sentences?|words)\b", re.I)


def strip_style(text: str) -> str:
    t = _STYLE_RX.sub(" ", text)
    t = re.sub(r"\s*,\s*(?=[,.?!]|$)", "", t)
    t = " ".join(t.split()).strip(" ,")
    return t if len(t.split()) >= 2 else text


def normalize_query(text: str, scheme: dict | None) -> str:
    text = strip_style(text)
    if not scheme or not scheme.get("prefixes") and not scheme.get("sku_cats"):
        return text
    return normalize_ids(normalize_skus(text, scheme.get("sku_prefix", ""), scheme.get("sku_cats", [])),
                         scheme.get("prefixes", {}), scheme.get("bare"))


# ───────────────────────────── library ─────────────────────────────

_SCHEMA = """
PRAGMA journal_mode=WAL;
CREATE TABLE IF NOT EXISTS docs(
  id INTEGER PRIMARY KEY, sha256 TEXT UNIQUE NOT NULL, file TEXT NOT NULL, path TEXT,
  title TEXT, short_title TEXT, authors TEXT, year TEXT, subject TEXT, kind TEXT, licence TEXT, source_url TEXT,
  pages INTEGER, chunks INTEGER, sentences INTEGER, bytes INTEGER, added REAL,
  embed_model TEXT, dim INTEGER, chunker_ver INTEGER, report TEXT);
CREATE TABLE IF NOT EXISTS chunks(
  id INTEGER PRIMARY KEY, doc_id INTEGER NOT NULL, ord INTEGER, kind TEXT, section TEXT, header TEXT, text TEXT,
  codes TEXT, ntok INTEGER, page INTEGER, page_label TEXT, page_end INTEGER, emb BLOB, emb_scale REAL);
CREATE INDEX IF NOT EXISTS chunks_doc ON chunks(doc_id, ord);
CREATE TABLE IF NOT EXISTS sentences(
  id INTEGER PRIMARY KEY, chunk_id INTEGER NOT NULL, ord INTEGER, page INTEGER, page_label TEXT,
  text TEXT, ntok INTEGER, emb BLOB, emb_scale REAL);
CREATE INDEX IF NOT EXISTS sentences_chunk ON sentences(chunk_id, ord);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(header, text, codes,
  tokenize='porter unicode61 remove_diacritics 2');
CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT);
"""


class _Progress:
    def __init__(self, fn: ProgressFn, file: str) -> None:
        self.fn, self.file = fn, file
        self.phase = None
        self.last = 0.0
        self.t_phase = time.perf_counter()

    def __call__(self, phase: str, done: int = 0, total: int = 0, force: bool = False, **kw) -> None:
        if not self.fn:
            return
        now = time.perf_counter()
        if phase != self.phase:
            self.phase, self.t_phase, force = phase, now, True
        if not force and now - self.last < 0.25:
            return
        self.last = now
        el = max(now - self.t_phase, 1e-6)
        per_s = done / el if done else 0.0
        ev = {"phase": phase, "file": self.file, "per_s": round(per_s, 1),
              "eta_s": round((total - done) / per_s, 1) if per_s and total else None, "ts": time.time()}
        frac = kw.pop("frac", None)  # work fraction when items differ in cost (embedding: wordpiece share)
        if frac:
            ev["eta_s"] = round(el * (1 - frac) / frac, 1)
        ev.update(kw)
        try:
            self.fn(ev)
        except Exception as e:  # noqa: BLE001
            log.warning("progress callback failed: %s", e)


class Library:
    def __init__(self, profile: str, path: Path | None = None) -> None:
        if not re.fullmatch(r"[a-z0-9_-]{1,32}", profile):
            raise ValueError(f"bad profile name {profile!r}")
        self.profile = profile
        MEMORY_DIR.mkdir(parents=True, exist_ok=True)
        self.path = Path(path or (MEMORY_DIR / f"{profile}_library.sqlite"))
        self._lock = threading.RLock()
        self._ingest_lock = threading.Lock()
        self.db = sqlite3.connect(str(self.path), check_same_thread=False)
        self.db.executescript(_SCHEMA)
        self.db.execute("INSERT OR IGNORE INTO meta VALUES ('schema_ver', ?)", (str(SCHEMA_VER),))
        if self.db.execute("SELECT count(*) FROM chunks_fts_config WHERE k='rank'").fetchone()[0] == 0:
            self.db.execute("INSERT INTO chunks_fts(chunks_fts, rank) VALUES ('rank', 'bm25(3.0, 1.0, 8.0)')")
        self.db.commit()
        self._load()

    # ---------- in-memory index ----------
    def _load(self) -> None:
        t0 = time.perf_counter()
        rows = self.db.execute("SELECT c.id, c.doc_id, c.emb, c.emb_scale FROM chunks c WHERE c.emb IS NOT NULL ORDER BY c.id").fetchall()
        docs = {r[0]: (r[0], r[1], r[2], r[3], r[4], Path(r[5] or "").suffix.lower() in (".xlsx", ".csv"))
                for r in self.db.execute("SELECT id, title, short_title, subject, kind, file FROM docs")}
        boost: dict[int, float] = {}
        papers = {d for d, r in docs.items() if (r[4] or "").lower() in ("paper", "article", "preprint")}
        if papers:
            for cid, did, sec, pg in self.db.execute(
                    "SELECT id, doc_id, section, page FROM chunks WHERE emb IS NOT NULL AND doc_id IN (%s)" % ",".join(map(str, papers))):
                if pg == 0 or re.search(r"\b(abstract|introduction|summary|overview)\b", sec or "", re.I):
                    boost[cid] = PAPER_FRONT_BOOST
        name_idx: dict[str, set] = defaultdict(set)
        for did, title, short, authors in self.db.execute("SELECT id, title, short_title, authors FROM docs"):
            words = set(_terms(f"{title or ''} {short or ''}"))
            for a in (authors or "").split(",")[:3]:
                if a.strip() and "et al" not in a:
                    words.add(a.strip().split()[-1].lower())
            for w in words:
                if len(w) >= 4 and w not in _GATE_IGNORE:
                    name_idx[w].add(did)
        name_idx = {w: ds for w, ds in name_idx.items() if len(ds) == 1} if len(docs) >= 3 else {}
        scheme = learn_id_scheme(r[0] for r in self.db.execute("SELECT text FROM chunks WHERE codes != '' AND kind = 'row'"))
        with self._lock:
            self._docs = docs
            self._boost = boost
            self._name_idx = name_idx
            self._summ = {r[0] for r in self.db.execute("SELECT id FROM chunks WHERE kind='summary'")}
            self._scheme = scheme
            self._cid = np.array([r[0] for r in rows], np.int64)
            self._did = np.array([r[1] for r in rows], np.int64)
            if rows:
                Q = np.frombuffer(b"".join(r[2] for r in rows), np.int8).reshape(len(rows), embedder.DIM)
                self._X = Q.astype(np.float32) * np.array([r[3] for r in rows], np.float32)[:, None]
            else:
                self._X = np.zeros((0, embedder.DIM), np.float32)
            self._pos = {int(c): i for i, c in enumerate(self._cid)}
            self._chunk_doc = {int(c): int(d) for c, d in zip(self._cid, self._did)}
            self._n_fts = self.db.execute("SELECT count(*) FROM chunks WHERE kind != 'refs'").fetchone()[0]
            self._dv = self.db.execute("PRAGMA data_version").fetchone()[0]
        log.info("library %s: %d docs, %d vectors loaded in %.0f ms", self.profile, len(docs), len(rows), (time.perf_counter() - t0) * 1e3)

    def _maybe_reload(self) -> None:
        """Pick up docs ingested by another process (e.g. rag_ingest.py while the server runs)."""
        with self._lock:
            dv = self.db.execute("PRAGMA data_version").fetchone()[0]
            if dv != getattr(self, "_dv", dv):
                self._load()
            self._dv = dv

    # ---------- ingest ----------
    def ingest_file(self, path, meta: dict | None = None, progress: ProgressFn = None) -> dict:
        path = Path(path)
        ext = path.suffix.lower()
        if ext not in SUPPORTED:
            raise ValueError(f"unsupported file type {ext}")
        t0 = time.perf_counter()
        why = self._skip_reason(path, meta)
        if why:
            self._drop_file(path.name)
            return {"file": path.name, "skipped": True, "reason": why}
        sha = _sha256(path)
        with self._lock:
            row = self.db.execute("SELECT id, title FROM docs WHERE sha256=?", (sha,)).fetchone()
        if row:
            m = dict(meta) if meta else find_manifest_meta(path)
            if m:  # the manifest may have arrived after the first ingest: refresh the catalogue fields
                with self._lock, self.db:
                    self.db.execute("UPDATE docs SET title=coalesce(?,title), short_title=coalesce(?,short_title), authors=coalesce(?,authors),"
                                    " year=coalesce(?,year), subject=coalesce(?,subject), kind=coalesce(?,kind), licence=coalesce(?,licence),"
                                    " source_url=coalesce(?,source_url) WHERE id=?",
                                    (m.get("title"), m.get("short_title"), _s(m.get("authors")), _s(m.get("year")), m.get("subject"),
                                     m.get("kind"), m.get("licence"), m.get("source_url"), row[0]))
                self._load()
            return {"file": path.name, "skipped": True, "reason": "already ingested", "doc_id": row[0], "title": m.get("title") or row[1]}
        with self._ingest_lock:
            return self._ingest(path, ext, sha, meta, progress, t0)

    def _ingest(self, path: Path, ext: str, sha: str, meta: dict | None, progress: ProgressFn, t0: float) -> dict:
        twin = _twin_pdf(path) if ext in (".md", ".markdown", ".txt") else None
        m = dict(meta) if meta else (find_manifest_meta(twin) if twin else {}) or find_manifest_meta(path)
        title = m.get("title") or path.stem.replace("_", " ")
        short = spoken_short_title(m, title, path.stem)
        front = bool(re.search(r"\b(?:prelims|front matter)\b", title, re.I)) and not re.search(r"\bCh(apter)?\.?\s*\d", str(m.get("chapter") or title))
        prog = _Progress(progress, path.name)
        ms: dict[str, float] = {}
        tok = _Tok()
        t = time.perf_counter()
        kind_default = "sheet" if ext in (".xlsx", ".csv") else ("notes" if ext in (".md", ".markdown", ".txt") else "book")
        doc_kind = m.get("kind") or kind_default

        if ext in (".xlsx", ".csv"):
            prog("extract", 0, 1, page=0, pages=1, chunks=0)
            sheets = read_sheets(path)
            ms["extract"] = (time.perf_counter() - t) * 1e3
            t = time.perf_counter()
            chunks = []
            for si, (name, rows) in enumerate(sheets.items()):
                for c in sheet_chunks(name, rows):
                    n = tok.counts([c["text"]])[0]
                    chunks.append({"kind": c["kind"], "section": name, "header": f"{short} › {name}",
                                   "sents": [(c["text"], si, c["label"], n)], "page": si, "page_label": c["label"], "page_end": si})
                prog("chunk", si + 1, len(sheets), page=si + 1, pages=len(sheets), chunks=len(chunks))
            ext_info = {"pages": len(sheets), "scanned": [], "headings": "sheets", "labels": list(sheets)}
        else:
            if ext == ".pdf":
                ext_info = extract_pdf(path, lambda i, n: prog("extract", i, n, page=i, pages=n, chunks=0))
            else:
                prog("extract", 0, 1, page=0, pages=1, chunks=0)
                ext_info = extract_text(path, markdown=ext in (".md", ".markdown"))
                mapped = _map_blocks_to_pdf(ext_info["blocks"], twin) if twin else None
                if mapped:  # cite the twin PDF's printed pages
                    ext_info["blocks"], ext_info["labels"] = mapped
                    ext_info["pages"] = len(mapped[1])
            ms["extract"] = (time.perf_counter() - t) * 1e3
            t = time.perf_counter()
            prog("chunk", 0, ext_info["pages"], page=ext_info["pages"], pages=ext_info["pages"], chunks=0, force=True)
            chunks = build_chunks(ext_info, short, tok)
            if front:  # front matter (prelims) / contents file: stored, not indexed
                for c in chunks:
                    c["kind"] = "refs"
        ms["chunk"] = (time.perf_counter() - t) * 1e3
        npages = ext_info["pages"]
        prog("chunk", npages, npages, page=npages, pages=npages, chunks=len(chunks), force=True)

        # embed: chunk = header + text; sentences ≥ SENT_MIN_WORDS words (rows: the chunk vector is reused)
        t = time.perf_counter()
        live = [c for c in chunks if c["kind"] != "refs"]
        ctexts = [c["header"] + "\n" + " ".join(s[0] for s in c["sents"]) for c in live]
        stexts, sidx = [], []
        for ci, c in enumerate(live):
            if len(c["sents"]) > 1:
                for si, s in enumerate(c["sents"]):
                    if len(s[0].split()) >= SENT_MIN_WORDS:
                        stexts.append(s[0])
                        sidx.append((ci, si))
        total = len(ctexts) + len(stexts)
        E = self._embed_all(ctexts + stexts, lambda d, f: prog("embed", d, total, page=npages, pages=npages, chunks=len(chunks), embedded=d, embed_total=total, frac=f))
        CV, SV = E[:len(ctexts)], E[len(ctexts):]
        svec: dict[tuple[int, int], np.ndarray] = {k: SV[i] for i, k in enumerate(sidx)}
        ms["embed"] = (time.perf_counter() - t) * 1e3

        # index (own connection, so searches on self.db keep running under WAL)
        t = time.perf_counter()
        prog("index", 0, len(chunks), page=npages, pages=npages, chunks=len(chunks), force=True)
        w = sqlite3.connect(str(self.path), timeout=30)
        nsent = 0
        try:
            with w:
                cur = w.execute("INSERT INTO docs(sha256, file, path, title, short_title, authors, year, subject, kind, licence, source_url,"
                                " pages, chunks, sentences, bytes, added, embed_model, dim, chunker_ver, report) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                                (sha, path.name, str(path), title, short, _s(m.get("authors")), _s(m.get("year")), m.get("subject"), doc_kind,
                                 m.get("licence"), m.get("source_url"), npages, len(chunks), 0, path.stat().st_size, time.time(),
                                 embedder.MODEL_NAME, embedder.DIM, CHUNKER_VER, None))
                did = cur.lastrowid
                Qc, Sc = _quant(CV) if len(CV) else (np.zeros((0, embedder.DIM), np.int8), np.zeros(0, np.float32))
                li = 0
                for o, c in enumerate(chunks):
                    text = " ".join(s[0] for s in c["sents"])
                    codes = "" if c["kind"] == "summary" else " ".join(sorted(set(_codes(text))))  # IDs listed in summaries must not outrank the row
                    emb, sc = (None, None)
                    if c["kind"] != "refs":
                        emb, sc = Qc[li].tobytes(), float(Sc[li])
                    cid = w.execute("INSERT INTO chunks(doc_id, ord, kind, section, header, text, codes, ntok, page, page_label, page_end, emb, emb_scale)"
                                    " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                                    (did, o, c["kind"], c["section"], c["header"], text, codes, sum(s[3] for s in c["sents"]),
                                     c["page"], c["page_label"], c["page_end"], emb, sc)).lastrowid
                    if c["kind"] != "refs":
                        w.execute("INSERT INTO chunks_fts(rowid, header, text, codes) VALUES (?,?,?,?)", (cid, c["header"], text, codes))
                        for si, (s, pg, lab, nt) in enumerate(c["sents"]):
                            v = svec.get((li, si))
                            if v is None and len(c["sents"]) == 1:
                                v = CV[li]
                            if v is not None:
                                q8, s8 = _quant(v[None, :])
                                eb, es = q8[0].tobytes(), float(s8[0])
                            else:
                                eb, es = None, None
                            w.execute("INSERT INTO sentences(chunk_id, ord, page, page_label, text, ntok, emb, emb_scale) VALUES (?,?,?,?,?,?,?,?)",
                                      (cid, si, pg, lab, s, nt, eb, es))
                            nsent += 1
                        li += 1
                    prog("index", o + 1, len(chunks), page=npages, pages=npages, chunks=len(chunks))
                ms["index"] = (time.perf_counter() - t) * 1e3
                secs = time.perf_counter() - t0
                report = {"ms_by_phase": {k: round(v) for k, v in ms.items()}, "headings_source": ext_info.get("headings"),
                          "scanned_pages": [ext_info["labels"][i] if i < len(ext_info["labels"]) else str(i + 1) for i in ext_info["scanned"]],
                          "refs_chunks": sum(c["kind"] == "refs" for c in chunks), "seconds": round(secs, 2)}
                w.execute("UPDATE docs SET sentences=?, report=? WHERE id=?", (nsent, json.dumps(report), did))
        finally:
            w.close()
        self._load()
        secs = time.perf_counter() - t0
        res = {"file": path.name, "skipped": False, "doc_id": did, "title": title, "short_title": short, "pages": npages,
               "chunks": len(chunks), "indexed_chunks": len(live), "sentences": nsent, "scanned_pages": len(ext_info["scanned"]),
               "headings": ext_info.get("headings"), "seconds": round(secs, 2),
               "pages_per_s": round(npages / secs, 1), "chunks_per_s": round(len(chunks) / secs, 1), "ms_by_phase": {k: round(v) for k, v in ms.items()}}
        prog("done", npages, npages, page=npages, pages=npages, chunks=len(chunks), force=True, report=res)
        log.info("ingested %s: %d pages, %d chunks, %d sentences in %.1f s", path.name, npages, len(chunks), nsent, secs)
        return res

    def skip_patterns(self) -> list[str]:
        with self._lock:
            r = self.db.execute("SELECT value FROM meta WHERE key='skip'").fetchone()
        return DEFAULT_SKIP.get(self.profile, []) + (json.loads(r[0]) if r else [])

    def _skip_reason(self, path: Path, meta: dict | None) -> str | None:
        import fnmatch
        for pat in self.skip_patterns():
            if fnmatch.fnmatch(path.name, pat) or fnmatch.fnmatch(str(path), pat):
                return f"excluded by skip list ({pat})"
        m = meta if meta is not None else find_manifest_meta(path)
        if m and (m.get("exclude") or m.get("skip") or m.get("index") is False):
            return "excluded by manifest flag"
        return None

    def _drop_file(self, name: str) -> int:
        with self._lock:
            ids = [r[0] for r in self.db.execute("SELECT id FROM docs WHERE file=?", (name,))]
        for i in ids:
            self.delete_doc(i)
        return len(ids)

    def exclude(self, pattern: str) -> int:
        """Exclude files matching a glob from this library (removes them now and skips them on later ingests)."""
        pats = self.skip_patterns()[len(DEFAULT_SKIP.get(self.profile, [])):]
        if pattern not in pats:
            pats.append(pattern)
        with self._lock, self.db:
            self.db.execute("INSERT OR REPLACE INTO meta VALUES ('skip', ?)", (json.dumps(pats),))
        import fnmatch
        with self._lock:
            names = [r[0] for r in self.db.execute("SELECT file FROM docs")]
        return sum(self._drop_file(n) for n in names if fnmatch.fnmatch(n, pattern))

    _GENERIC = {"stock", "stocks", "inventory", "many", "much", "total", "units", "unit", "left", "available", "have", "sold",
                "order", "orders", "price", "cost", "piece", "pieces", "item", "items", "product", "products", "sku"}

    def _named_row(self, summary_id: int, terms: list[str]) -> int | None:
        """A category summary lists its products ("4 product names: A, B, C, D"). When exactly one of them carries every
        specific word of the question, return that product's row in the same sheet."""
        with self._lock:
            r = self.db.execute("SELECT doc_id, section, text FROM chunks WHERE id=?", (summary_id,)).fetchone()
        if not r:
            return None
        m = re.search(r"product names?: ([^)]*)\)", r[2] or "")
        q = {_stem5(w) for w in terms if w not in self._GENERIC}
        if not m or not q:
            return None
        hits = [n.strip() for n in m.group(1).split(",") if q <= {_stem5(x) for x in _terms(n)}]
        if len(hits) != 1:
            return None
        with self._lock:
            row = self.db.execute("SELECT id FROM chunks WHERE doc_id=? AND section=? AND kind='row' AND text LIKE ? ORDER BY ord LIMIT 1",
                                  (r[0], r[1], f"%product name: {hits[0]} ·%")).fetchone()
        return row[0] if row else None

    def _embed_all(self, texts: list[str], cb) -> np.ndarray:
        if not texts:
            return np.zeros((0, embedder.DIM), np.float32)
        emb = embedder.get()
        order = sorted(range(len(texts)), key=lambda i: len(texts[i]))  # length-sorted batches: little padding
        out = np.zeros((len(texts), embedder.DIM), np.float32)
        B = 32
        cost = [min(len(t), 1200) + 40 for t in texts]  # ≈ wordpieces (truncated at 256) + per-item overhead
        tot, acc = float(sum(cost)), 0.0
        for b in range(0, len(order), B):
            idx = order[b:b + B]
            out[idx] = emb.encode([texts[i] for i in idx])
            acc += sum(cost[i] for i in idx)
            cb(min(b + B, len(order)), acc / tot)
        return out

    def ingest_dir(self, d, progress: ProgressFn = None) -> dict:
        d = Path(d)
        t0 = time.perf_counter()
        files = sorted(p for p in d.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED
                       and not p.name.startswith(".") and not p.stem.lower().startswith("readme"))
        # prefer an .xlsx over per-sheet .csv exports of the same sheet
        xl_sheets: dict[Path, set[str]] = defaultdict(set)
        for p in files:
            if p.suffix.lower() == ".xlsx":
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(str(p), read_only=True)
                    names = {re.sub(r"[^a-z0-9]", "", s.lower()) for s in wb.sheetnames}
                    wb.close()
                    xl_sheets[p.parent] |= names | {re.sub(r"[^a-z0-9]", "", (p.stem + s).lower()) for s in wb.sheetnames}
                except Exception as e:  # noqa: BLE001
                    log.warning("cannot read sheet names of %s: %s", p.name, e)
        res = {"docs": 0, "skipped": 0, "deduped_csv": [], "deduped_twins": [], "pages": 0, "chunks": 0, "scanned_pages": 0, "errors": [], "files": []}
        stems_md = {(p.parent, p.stem.lower()) for p in files if p.suffix.lower() in (".md", ".markdown", ".txt")}
        for p in sorted(files, key=lambda p: (p.suffix.lower() == ".csv", str(p))):
            if p.suffix.lower() == ".pdf" and (p.parent, p.stem.lower()) in stems_md:
                res["deduped_twins"].append(p.name)  # same document as the .md next to it (clean text; pages mapped from this PDF)
                self._drop_file(p.name)
                continue
            if p.suffix.lower() == ".csv" and p.parent in xl_sheets:
                k = re.sub(r"[^a-z0-9]", "", p.stem.lower())
                if k in xl_sheets[p.parent] or any(k.endswith(s) for s in xl_sheets[p.parent] if len(s) >= 4):
                    res["deduped_csv"].append(p.name)
                    continue
            try:
                r = self.ingest_file(p, progress=progress)
            except Exception as e:  # noqa: BLE001
                log.exception("ingest failed for %s", p)
                res["errors"].append({"file": p.name, "error": str(e)})
                continue
            res["files"].append(r)
            if r.get("skipped"):
                res["skipped"] += 1
            else:
                res["docs"] += 1
                res["pages"] += r["pages"]
                res["chunks"] += r["chunks"]
                res["scanned_pages"] += r["scanned_pages"]
        res["seconds"] = round(time.perf_counter() - t0, 2)
        return res

    # ---------- search ----------
    def search(self, query: str, k: int = 3, budget_tokens: int = 160, subject: str | None = None) -> dict:
        """Hybrid search → compressed context + gate. See module docstring; result keys are the integration contract."""
        T0 = time.perf_counter()
        ms: dict[str, float] = {}
        self._maybe_reload()
        empty = {"passages": [], "context": "", "top_score": 0.0, "gate": False, "hits": {"fts": 0, "vec": 0}, "query": query}
        with self._lock:
            X, cid_arr, did_arr, docs, _nfts, boost, scheme, summ = (self._X, self._cid, self._did, self._docs, self._n_fts, self._boost,
                                                                    self._scheme, self._summ)
            name_idx, c2d = self._name_idx, self._chunk_doc
        if not len(cid_arr) or not query.strip():
            return {**empty, "ms": {"fts": 0, "vec": 0, "fuse": 0, "compress": 0, "total": 0}}
        q_text = normalize_query(query, scheme)
        q_text = re.sub(r"\b(?:[A-Za-z]\s){1,5}[A-Za-z]\b", lambda m: m.group(0).replace(" ", "") if len(m.group(0)) >= 3 else m.group(0), q_text)  # "e s m" → "esm"  # spoken IDs/SKUs → canonical ("o r d ten forty two" → "ORD-1042")
        allowed = None
        if subject:
            allowed = {d for d, r in docs.items() if (r[3] or "").lower() == subject.lower()}

        # query embedding runs on a worker thread (ORT releases the GIL) while FTS5 runs here
        t_emb = time.perf_counter()
        fut = _pool().submit(lambda: (embedder.get().encode([q_text])[0], time.perf_counter()))

        # FTS5 BM25 (rank = bm25 with column weights header 3, text 1, codes 8)
        t = time.perf_counter()
        terms = _terms(q_text)
        if any(_latin(w)[-1:] and _latin(w)[-1].endswith("eus") for w in terms):
            terms = [w for w in terms if w not in ("muscle", "muscles")]  # "X muscle" may appear under the Latin term alone
        qcodes = _query_codes(q_text)
        extra = [x for w in terms for x in _EXPAND.get(w, []) + _latin(w) if x not in terms]
        fts: list[int] = []
        code_hits: list[int] = []
        if allowed is None:
            sql = "SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ? ORDER BY rank LIMIT ?"
        else:
            sql = ("SELECT f.rowid FROM chunks_fts f JOIN chunks c ON c.id=f.rowid WHERE chunks_fts MATCH ? AND c.doc_id IN (%s) ORDER BY f.rank LIMIT ?"
                   % (",".join(str(int(x)) for x in allowed) or "-1"))
        df: dict[str, int] = {}
        with self._lock:
            if qcodes:
                code_hits = [r[0] for r in self.db.execute(sql, (" OR ".join(f'codes:"{c}"' for c in qcodes), 10))]
            q_terms = terms + list(dict.fromkeys(extra)) + [c for c in qcodes if c not in terms]
            if q_terms:
                fts = [r[0] for r in self.db.execute(sql, (" OR ".join('"%s"' % w.replace('"', "") for w in q_terms), CAND))]
            for w in terms:  # library document frequency (capped): 0 = the word (or its stem prefix) is nowhere in the notes
                pre = w[:max(5, len(w) - 3)]
                expr = '"%s" OR %s*' % (w, pre) if pre.isalnum() else '"%s"' % w.replace('"', "")
                df[w] = self.db.execute("SELECT count(*) FROM (SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ? LIMIT 50)",
                                        (expr,)).fetchone()[0]
        ms["fts"] = (time.perf_counter() - t) * 1e3

        # dense: exact scan of the in-RAM chunk matrix
        q, t_done = fut.result()
        ms["embed"] = (t_done - t_emb) * 1e3
        t = time.perf_counter()
        sims = X @ q
        if allowed is not None:
            sims = np.where(np.isin(did_arr, list(allowed)), sims, -1.0)
        kk = min(CAND, len(sims))
        top = np.argpartition(-sims, kk - 1)[:kk]
        top = top[np.argsort(-sims[top])]
        vec = [int(cid_arr[i]) for i in top if sims[i] > -1.0]
        ms["vec"] = (time.perf_counter() - t) * 1e3

        # RRF fusion (exact ID hits count as an extra, double-weight list); small prior for paper abstracts/intros
        t = time.perf_counter()
        rrf: dict[int, float] = defaultdict(float)
        for r, c in enumerate(fts):
            rrf[c] += 1.0 / (RRF_K + r + 1)
        for r, c in enumerate(vec):
            rrf[c] += 1.0 / (RRF_K + r + 1)
        for r, c in enumerate(code_hits):
            rrf[c] += 2.0 / (RRF_K + r + 1)
        agg = bool(_AGG_Q_RX.search(q_text))
        named = set()
        for w in terms:  # the question names a document ("how does <author>…", "<a story title>") with a word rare in the library
            if w in name_idx:
                with self._lock:  # rare in the body text (headers carry the title on every chunk, so they don't count)
                    n_txt = self.db.execute("SELECT count(*) FROM (SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ? LIMIT 60)",
                                            ('text : "%s"' % w,)).fetchone()[0]
                if n_txt <= GATE_RARE_DF * 4:
                    named |= name_idx[w]
        if named and len(named) <= 2:
            for c in list(rrf):
                if c2d.get(c) in named:
                    rrf[c] *= DOC_MENTION_BOOST
        sched_docs = {d for d, r in docs.items() if (r[4] or "").lower() in SCHEDULE_KINDS} if _SCHEDULE_Q_RX.search(q_text) else set()
        if sched_docs:  # make sure the diary's chunks are candidates even when its words are common in the textbooks
            with self._lock:
                for r in self.db.execute("SELECT f.rowid FROM chunks_fts f JOIN chunks c ON c.id=f.rowid WHERE chunks_fts MATCH ? AND c.doc_id IN (%s) ORDER BY f.rank LIMIT 10"
                                         % ",".join(map(str, sched_docs)), (" OR ".join('"%s"' % w for w in terms) or '""',)):
                    rrf[r[0]] += 1.0 / (RRF_K + 1)
            chunk_doc = c2d
        for c in rrf:
            rrf[c] *= boost.get(c, 1.0) * (SUMMARY_BOOST if agg and c in summ else 1.0)
            if sched_docs and chunk_doc.get(c) in sched_docs:
                rrf[c] *= SCHEDULE_BOOST
        if agg:  # "which city has the most orders": the summary grouped BY the asked column wins
            qs = {_stem5(w) for w in terms}
            with self._lock:  # summaries grouped by a column the question names may not be BM25/vector candidates at all
                for w in terms:
                    for (cid,) in self.db.execute("SELECT id FROM chunks WHERE kind='summary' AND (text LIKE ? OR text LIKE ?) LIMIT 3",
                                                  (f"% by {w} (%", f"% by {w}s (%")):
                        if cid not in rrf:
                            rrf[cid] = max(rrf.values(), default=1.0 / RRF_K) * 0.9
            cand = [c for c in rrf if c in summ]
            if cand:
                with self._lock:
                    st = dict(self.db.execute("SELECT id, text FROM chunks WHERE id IN (%s)" % ",".join(map(str, cand))).fetchall())
                for c in cand:
                    m = re.match(r"[\w ]+? by (?:month of )?([a-z][\w ]*?) \(", st.get(c) or "")
                    if m and all(_stem5(w) in qs for w in _terms(m.group(1))):
                        rrf[c] *= GROUP_BY_BOOST
        ranked = sorted(rrf, key=lambda c: -rrf[c])[:max(k, 1)]
        row = None
        if ranked and ranked[0] in summ:
            row = self._named_row(ranked[0], terms)
            if row is not None:  # "how many yellow rain jackets": that product's own row leads, the category total follows
                # the category total is left out: a 2B model given both "total stock" numbers picks the wrong one
                ranked = [row] + [c for c in sorted(rrf, key=lambda c: -rrf[c]) if c != row and c not in summ][: max(k, 1) - 1]
        ms["fuse"] = (time.perf_counter() - t) * 1e3
        if not ranked:
            ms["total"] = (time.perf_counter() - T0) * 1e3
            return {**empty, "query_norm": q_text, "ms": {k2: round(v, 2) for k2, v in ms.items()}, "hits": {"fts": len(fts), "vec": len(vec)}}

        # compression: sentence windows scored by 0.55·cos + 0.25·rrf + 0.20·term overlap
        t = time.perf_counter()
        ph = ",".join("?" * len(ranked))
        with self._lock:
            crow = {r[0]: r for r in self.db.execute(
                f"SELECT id, doc_id, section, page, page_label, kind, text, header FROM chunks WHERE id IN ({ph})", ranked)}
            srows = self.db.execute(
                f"SELECT id, chunk_id, ord, page, page_label, text, ntok, emb, emb_scale FROM sentences WHERE chunk_id IN ({ph}) ORDER BY chunk_id, ord", ranked).fetchall()
        stems = {_stem5(w) for w in terms}
        rmax = max(rrf[c] for c in ranked)
        S = []
        for r in srows:
            cos = float((np.frombuffer(r[7], np.int8).astype(np.float32) * r[8]) @ q) if r[7] is not None else 0.0
            words = {_stem5(w) for w in _terms(r[5])}
            ov = len(stems & words) / max(len(stems), 1)
            code_bonus = 0.3 if qcodes and any(c in "".join(_codes(r[5])) for c in qcodes) else 0.0
            S.append({"id": r[0], "chunk": r[1], "ord": r[2], "page": r[3], "label": r[4], "text": r[5], "ntok": r[6],
                      "cos": cos, "score": 0.55 * cos + 0.25 * rrf[r[1]] / rmax + 0.20 * ov + code_bonus})
        by_chunk: dict[int, dict[int, dict]] = defaultdict(dict)
        for s in S:
            by_chunk[s["chunk"]][s["ord"]] = s
        tk = embedder.get().tok

        def ntok(text: str) -> int:
            return len(tk.encode(text, add_special_tokens=False).ids)

        def prefix(c: int, s: dict) -> str:
            cr = crow[c]
            d = docs.get(cr[1])
            short = (d[2] or d[1]) if d else "notes"
            sec = (cr[2] or "").split(" › ")[-1]
            if d and d[5]:  # spreadsheet: "[Orders workbook, Orders sheet, row 42]"
                lab = s["label"] or ""
                m = re.match(r"(.+) row (\d+)$", lab)
                loc = f"{m.group(1)} sheet, row {m.group(2)}" if m else (lab.replace(" summary", " sheet summary") if lab.endswith(" summary") else lab)
                return f"[{short}, {loc}]"
            loc = f", p. {s['label']}" if s["label"] else ""
            return f"[{short}{loc}" + (f", {sec}" if sec and sec != short else "") + "]"

        chosen: set[int] = set()
        used = 0
        heads: dict[int, int] = {}
        best_score = max((s["score"] for s in S), default=0.0)
        first = max(by_chunk[ranked[0]].values(), key=lambda s: s["score"]) if by_chunk.get(ranked[0]) else None
        top_s = max(S, key=lambda s: s["score"]) if S else None
        if row is None and ranked[0] not in summ and first is not None and top_s is not None and top_s["score"] > first["score"] + LEAD_MARGIN:
            first = top_s  # a clearly better sentence in another passage leads (a long, loosely related top passage must not fill the budget)
        order_s = ([first] if first else []) + sorted((s for s in S if s is not first), key=lambda s: -s["score"])
        truncated: dict[int, str] = {}
        if first is not None and first["chunk"] in summ:
            pre = prefix(first["chunk"], first)
            if ntok(pre) + first["ntok"] + 2 > budget_tokens:
                # a long grouped summary ("Orders by city: <city> 11, …" for many cities) is sorted biggest first:
                # keep its head within ~70 % of the budget instead of dropping the one passage that answers it
                words = first["text"].split()
                while words and ntok(pre + " " + " ".join(words) + " …") > int(budget_tokens * 0.7):
                    words = words[: max(1, int(len(words) * 0.85))] if len(words) > 1 else []
                if words:
                    truncated[first["id"]] = " ".join(words) + " …"
                    chosen.add(first["id"])
                    heads[first["chunk"]] = 0
                    used = ntok(pre + " " + truncated[first["id"]]) + 1
        for s in order_s:
            if s is not first and s["score"] < 0.5 * best_score:  # don't pad the budget with weak sentences
                break
            win = [s] + [by_chunk[s["chunk"]].get(s["ord"] + d) for d in WINDOW]
            for j, x in enumerate(win):
                if x is None or x["id"] in chosen:
                    continue
                need = x["ntok"] + 1
                if x["chunk"] not in heads:
                    need += ntok(prefix(x["chunk"], x)) + 1
                if used + need > budget_tokens:
                    if j == 0:
                        break
                    continue
                if x["chunk"] not in heads:
                    heads[x["chunk"]] = len(heads)
                chosen.add(x["id"])
                used += need
            if used >= budget_tokens - 8:
                break
        if not chosen and S:  # best sentence alone is over budget (long table row): cut it to fit rather than return nothing
            s = max(S, key=lambda s: s["score"])
            pre = prefix(s["chunk"], s)
            words = s["text"].split()
            while words and ntok(pre + " " + " ".join(words) + " …") > budget_tokens:
                words = words[: max(1, int(len(words) * 0.85))] if len(words) > 1 else []
            if words:
                truncated[s["id"]] = " ".join(words) + " …"
                chosen.add(s["id"])
                used = ntok(pre + " " + truncated[s["id"]])
        passages, blocks = [], []
        for c in ranked:
            cr = crow.get(c)
            if cr is None:
                continue
            d = docs.get(cr[1])
            sel = [s for s in sorted(by_chunk[c].values(), key=lambda s: s["ord"]) if s["id"] in chosen]
            best = max(by_chunk[c].values(), key=lambda s: s["score"]) if by_chunk[c] else None
            ref = sel[0] if sel else best
            text = " ".join(truncated.get(s["id"], s["text"]) for s in sel)
            if sel:
                blocks.append(prefix(c, ref) + " " + text)
            passages.append({"doc_id": cr[1], "chunk_id": c, "title": d[1] if d else "", "short_title": (d[2] if d else ""),
                             "subject": d[3] if d else None, "kind": d[4] if d else None,
                             "page_label": ref["label"] if ref else cr[4], "page": (ref["page"] if ref else cr[3]) + 1,
                             "cite": prefix(c, ref).strip("[]") if ref else "",
                             "section": cr[2], "text": text or (best["text"] if best else cr[6][:300]), "in_context": bool(sel),
                             "score": round(max((s["cos"] for s in by_chunk[c].values()), default=0.0), 3), "rrf": round(rrf[c], 4)})
        context = "\n".join(blocks)
        ms["compress"] = (time.perf_counter() - t) * 1e3

        # gate: dense evidence AND lexical evidence (see GATE_* constants)
        top_cos = max((s["cos"] for s in S), default=0.0)
        content = [w for w in terms if len(w) >= 3 and w not in _GATE_IGNORE]
        cover = 0.0
        if content:
            for c in ranked:
                cr = crow.get(c)
                if cr is None:
                    continue
                cw = {_stem5(w) for w in _terms((cr[7] or "") + " " + (cr[6] or ""))}
                cover = max(cover, sum(any(_stem5(x) in cw for x in [w] + _EXPAND.get(w, [])) for w in content) / len(content))
        oov = [w for w in content if df.get(w, 1) == 0 and len(w) >= 3]
        top_doc = docs.get(crow[ranked[0]][1]) if ranked[0] in crow else None
        lenient = bool(top_doc and (top_doc[5] or (top_doc[4] or "").lower() in LENIENT_KINDS))
        tau = GATE_COS - (GATE_LENIENT_DELTA if lenient else 0.0)
        code_ok = bool(code_hits and code_hits[0] == ranked[0])
        top_words = set()
        for c in ranked[:2]:
            if c in crow:
                top_words |= {_stem5(w) for w in _terms((crow[c][7] or "") + " " + (crow[c][6] or ""))}
        rare_hit = any(0 < df.get(w, 0) <= GATE_RARE_DF and _stem5(w) in top_words for w in content)
        if not lenient and _RECENT_Q_RX.search(q_text):
            tau += GATE_STRONG  # recent events are not in books
        if code_ok:
            gate, why = True, "exact id"
        elif len(content) < GATE_MIN_WORDS and not rare_hit and not (named and c2d.get(ranked[0]) in named):
            gate, why = False, "too few content words"
        elif oov:
            gate, why = top_cos >= tau + GATE_OOV_EXTRA and cover >= 0.5, "oov:" + ",".join(oov)
        elif top_cos >= tau + GATE_STRONG:
            gate, why = cover >= 0.34 or len(content) <= 1, "strong cos"
        elif top_cos >= tau:
            gate, why = cover >= GATE_COVER, "cos+keywords"
        elif top_cos >= tau - GATE_LEX_BAND:
            gate, why = cover >= 0.99 and len(content) >= 2, "all keywords"
        else:
            gate, why = False, "weak"
        ms["total"] = (time.perf_counter() - T0) * 1e3
        return {"passages": passages, "context": context, "top_score": round(top_cos, 3), "gate": bool(gate),
                "gate_info": {"reason": why, "cover": round(cover, 2), "oov": oov, "tau": round(tau, 3)},
                "query": query, "query_norm": q_text,
                "ms": {k2: round(v, 2) for k2, v in ms.items()}, "hits": {"fts": len(fts), "vec": len(vec), "code": len(code_hits)},
                "tokens": used}

    # ---------- info ----------
    def docs(self) -> list[dict]:
        with self._lock:
            rows = self.db.execute("SELECT id, title, short_title, kind, pages, chunks, licence, subject, file, sentences, report FROM docs ORDER BY id").fetchall()
        return [{"doc_id": r[0], "title": r[1], "short_title": r[2], "kind": r[3], "pages": r[4], "chunks": r[5], "licence": r[6],
                 "subject": r[7], "file": r[8], "sentences": r[9], "report": json.loads(r[10]) if r[10] else None} for r in rows]

    def stats(self) -> dict:
        with self._lock:
            nd, np_, nc = self.db.execute("SELECT count(*), coalesce(sum(pages),0), coalesce(sum(chunks),0) FROM docs").fetchone()
            ns = self.db.execute("SELECT count(*) FROM sentences").fetchone()[0]
            n = len(self._cid)
        disk = sum(p.stat().st_size for p in self.path.parent.glob(self.path.name + "*"))
        return {"profile": self.profile, "path": str(self.path), "docs": nd, "pages": np_, "chunks": nc, "indexed_chunks": n,
                "sentences": ns, "vector_ram_mb": round(self._X.nbytes / 2**20, 2), "disk_mb": round(disk / 2**20, 2),
                "embed_model": embedder.MODEL_NAME, "gate_cos": GATE_COS}

    def delete_doc(self, doc_id: int) -> bool:
        with self._lock, self.db:
            ids = [r[0] for r in self.db.execute("SELECT id FROM chunks WHERE doc_id=?", (doc_id,))]
            for c in ids:
                self.db.execute("DELETE FROM chunks_fts WHERE rowid=?", (c,))
            self.db.execute("DELETE FROM sentences WHERE chunk_id IN (SELECT id FROM chunks WHERE doc_id=?)", (doc_id,))
            self.db.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
            n = self.db.execute("DELETE FROM docs WHERE id=?", (doc_id,)).rowcount
        self._load()
        return bool(n)

    def clear(self) -> None:
        """Remove every document (keeps the skip list). Used by `rag_ingest.py --rebuild`."""
        with self._lock, self.db:
            for t in ("sentences", "chunks", "docs"):
                self.db.execute(f"DELETE FROM {t}")
            self.db.execute("DELETE FROM chunks_fts")
        self._load()

    def close(self) -> None:
        with self._lock:
            self.db.close()


def _s(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (list, tuple)):
        return ", ".join(map(str, v))
    return str(v)


_executor = None


def _pool():
    global _executor
    if _executor is None:
        from concurrent.futures import ThreadPoolExecutor
        _executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="rag-embed")
    return _executor


def reload(profile: str) -> Library:
    """Force the cached Library to re-read its DB (searches also auto-reload when another process commits)."""
    lib = get(profile)
    lib._load()
    return lib


_libs: dict[str, Library] = {}
_libs_lock = threading.Lock()


def get(profile: str) -> Library:
    with _libs_lock:
        lib = _libs.get(profile)
        if lib is None:
            lib = _libs[profile] = Library(profile)
        return lib
